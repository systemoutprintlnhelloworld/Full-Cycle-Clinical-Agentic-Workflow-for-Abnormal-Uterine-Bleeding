"""Build metrics registry (audit-friendly machine metrics).

This registry is meant to answer:
- 每个机器指标（metric）对应哪个可审计的数据源文件？
- 数据源是否遵守样本排除规则（D1 anomaly / Gate3->D4）？
- 明细 / 汇总分别在哪个 sheet？样本量是多少？

Outputs:
- analysis_viz/docs/metrics_registry.csv

Notes:
- 当前只纳入已完成“明细+汇总+排除规则”结构化的指标源表。
- 论文 figdata 的大量 *_merged_source.xlsx 主要是“汇总”级别，
  它们的明细与排除审计需要单独补齐（后续会在 docs/todo.md 跟踪）。
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
ANALYSIS_VIZ = ROOT / "analysis_viz"

METRICS_REGISTRY = ANALYSIS_VIZ / "docs" / "metrics_registry.csv"


@dataclass(frozen=True)
class MetricDef:
    metric_id: str
    metric_name: str
    workbook_relpath: str
    detail_sheets: tuple[str, ...]
    summary_sheets: tuple[str, ...]
    d4_stage_values: tuple[str, ...] = ()
    d4_column_name: str | None = None
    gate3_excluded_sheet: str | None = None


def _rel(p: Path) -> str:
    try:
        return str(p.relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(p).replace("\\", "/")


def _read_meta_kv(wb: openpyxl.Workbook) -> dict[str, str]:
    if "meta" not in wb.sheetnames:
        return {}
    ws = wb["meta"]
    kv: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        k = row[0]
        v = row[1] if len(row) > 1 else None
        if k is None:
            continue
        kv[str(k).strip()] = "" if v is None else str(v).strip()
    return kv


def _sheet_case_ids(wb: openpyxl.Workbook, sheet_name: str) -> set[str]:
    if sheet_name not in wb.sheetnames:
        return set()
    ws = wb[sheet_name]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    if not header:
        return set()
    idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
    if "case_id" not in idx:
        return set()
    out: set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = row[idx["case_id"]]
        if cid is None:
            continue
        out.add(str(cid).strip())
    return out


def _sheet_keys(wb: openpyxl.Workbook, sheet_name: str) -> set[tuple[str, str, str]]:
    """Return (center, model, case_id) keys for audit checks.

    Note: 对 Sankey 这类“每个 case_id 会在多个 model 下重复出现”的明细表，
    仅使用 case_id 会产生误报，因此此处必须用三元组 key。
    """

    if sheet_name not in wb.sheetnames:
        return set()
    ws = wb[sheet_name]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    if not header:
        return set()
    idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
    if not all(k in idx for k in ("center", "model", "case_id")):
        return set()
    out: set[tuple[str, str, str]] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        center = row[idx["center"]]
        model = row[idx["model"]]
        cid = row[idx["case_id"]]
        if center is None or model is None or cid is None:
            continue
        out.add((str(center).strip(), str(model).strip(), str(cid).strip()))
    return out


def _sheet_has_any_true(wb: openpyxl.Workbook, sheet_name: str, col: str) -> bool:
    if sheet_name not in wb.sheetnames:
        return False
    ws = wb[sheet_name]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    if not header:
        return False
    idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
    if col not in idx:
        return False
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = row[idx[col]]
        if bool(v) is True:
            return True
    return False


def _gate3_d4_violation_by_stage(
    wb: openpyxl.Workbook,
    detail_sheets: tuple[str, ...],
    d4_stage_values: tuple[str, ...],
) -> bool:
    if not d4_stage_values:
        return False

    for sheet in detail_sheets:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        if not header:
            continue
        idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
        if "stage" not in idx or "is_gate3_fail" not in idx:
            continue

        for row in ws.iter_rows(min_row=2, values_only=True):
            stage = row[idx["stage"]]
            is_fail = row[idx["is_gate3_fail"]]
            if not is_fail:
                continue
            if stage is None:
                continue
            if str(stage).strip() in d4_stage_values:
                return True

    return False


def _gate3_d4_violation_by_blank_column(
    wb: openpyxl.Workbook,
    detail_sheet: str,
    excluded_sheet: str,
    d4_column_name: str,
) -> bool:
    if detail_sheet not in wb.sheetnames:
        return False
    if excluded_sheet not in wb.sheetnames:
        return False

    excluded_keys = _sheet_keys(wb, excluded_sheet)
    if not excluded_keys:
        return False

    ws = wb[detail_sheet]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    if not header:
        return False
    idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
    if not all(k in idx for k in ("center", "model", "case_id", d4_column_name)):
        return False

    for row in ws.iter_rows(min_row=2, values_only=True):
        center = row[idx["center"]]
        model = row[idx["model"]]
        cid = row[idx["case_id"]]
        if center is None or model is None or cid is None:
            continue
        key = (str(center).strip(), str(model).strip(), str(cid).strip())
        if key not in excluded_keys:
            continue
        v = row[idx[d4_column_name]]
        if v is None:
            continue
        if str(v).strip() != "":
            return True
    return False


def build_metrics_registry() -> None:
    metric_defs = [
        MetricDef(
            metric_id="llm_reasoning",
            metric_name="LLM 推理质量（reasoning）",
            workbook_relpath="analysis_viz/data/derived/metrics/llm_reasoning_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Rehab_Plan",),
        ),
        MetricDef(
            metric_id="llm_consistency",
            metric_name="LLM 一致性（consistency）",
            workbook_relpath="analysis_viz/data/derived/metrics/llm_consistency_source.xlsx",
            detail_sheets=("fact_detail_used", "cross_detail_used"),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4",),
        ),
        MetricDef(
            metric_id="llm_memory",
            metric_name="LLM 记忆（memory）",
            workbook_relpath="analysis_viz/data/derived/metrics/llm_memory_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Rehab_Plan",),
        ),
        MetricDef(
            metric_id="Fig7_sankey_flow",
            metric_name="Fig7 Sankey 流向（可解释明细+节点/连线汇总）",
            workbook_relpath="analysis_viz/data/derived/figdata/paper/Fig7__sankey_flow_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("nodes_summary_used", "links_summary_used"),
            d4_column_name="随访计划",
            gate3_excluded_sheet="excluded_d3_fail_d4_blank",
        ),
        MetricDef(
            metric_id="Fig3_check_match_rate",
            metric_name="Fig3 检查匹配度/合理性（可审计重建）",
            workbook_relpath="analysis_viz/data/derived/figdata/paper/Fig3__check_match_rate_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
        ),
        MetricDef(
            metric_id="Fig3_loop_inefficiency",
            metric_name="Fig3 无效循环率（0-match）/检查环节效率（可审计重建）",
            workbook_relpath="analysis_viz/data/derived/figdata/paper/Fig3__loop_inefficiency_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
        ),
        MetricDef(
            metric_id="algorithmic_check_match_rate",
            metric_name="Algorithmic A1 检查匹配率（病例均值）",
            workbook_relpath="analysis_viz/data/derived/metrics/algorithmic/algorithmic_check_match_rate_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
        ),
        MetricDef(
            metric_id="algorithmic_loop_inefficiency",
            metric_name="Algorithmic A2 无效循环率（0-match）",
            workbook_relpath="analysis_viz/data/derived/metrics/algorithmic/algorithmic_loop_inefficiency_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
        ),
        MetricDef(
            metric_id="algorithmic_stage_pass_rate",
            metric_name="Algorithmic A3 阶段通过率（D1/D2/D3）",
            workbook_relpath="analysis_viz/data/derived/metrics/algorithmic/algorithmic_stage_pass_rate_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
        ),
        MetricDef(
            metric_id="algorithmic_judge_scores_by_stage",
            metric_name="Algorithmic A4 Judge 综合评分（按阶段）",
            workbook_relpath="analysis_viz/data/derived/metrics/algorithmic/algorithmic_judge_scores_by_stage_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="algorithmic_judge_score_pathways",
            metric_name="Algorithmic A5 Judge score pathways（Diagnosis/Check/Plan）",
            workbook_relpath="analysis_viz/data/derived/metrics/algorithmic/algorithmic_judge_score_pathways_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="algorithmic_judge_scores_stagewise",
            metric_name="Algorithmic A6 6阶段 Judge Overall（Loop + Decision）",
            workbook_relpath="analysis_viz/data/derived/metrics/algorithmic/algorithmic_judge_scores_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="algorithmic_loop_inefficiency_stagewise",
            metric_name="Algorithmic A7 Loop无效率按阶段（D1/D2）",
            workbook_relpath="analysis_viz/data/derived/metrics/algorithmic/algorithmic_loop_inefficiency_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
        ),
        MetricDef(
            metric_id="algorithmic_stage_no_exit_rate",
            metric_name="Algorithmic A8 六阶段未发生退出比例",
            workbook_relpath="analysis_viz/data/derived/metrics/algorithmic/algorithmic_stage_no_exit_rate_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="calibration_ece_stagewise",
            metric_name="Calibration ECE stagewise（可审计）",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_ece_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="calibration_reliability_stagewise",
            metric_name="Calibration reliability stagewise?overall?",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_reliability_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used", "bin_summary_used"),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="calibration_bubble_check_stagewise",
            metric_name="Calibration bubble check stagewise",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_bubble_check_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used", "bin_summary_used"),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="calibration_bubble_diagnosis_stagewise",
            metric_name="Calibration bubble diagnosis stagewise",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_bubble_diagnosis_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used", "bin_summary_used"),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="calibration_bubble_plan_stagewise",
            metric_name="Calibration bubble plan stagewise",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_bubble_plan_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used", "bin_summary_used"),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="calibration_line_overall_stagewise",
            metric_name="Calibration line overall stagewise",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_line_overall_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used", "bin_summary_used"),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="calibration_line_check_stagewise",
            metric_name="Calibration line check stagewise",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_line_check_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used", "bin_summary_used"),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="calibration_line_diagnosis_stagewise",
            metric_name="Calibration line diagnosis stagewise",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_line_diagnosis_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used", "bin_summary_used"),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="calibration_line_plan_stagewise",
            metric_name="Calibration line plan stagewise",
            workbook_relpath="analysis_viz/data/derived/metrics/calibration/calibration_line_plan_stagewise_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used", "bin_summary_used"),
            d4_stage_values=("D4_Plan", "D4_Rehab_Plan"),
        ),
        MetricDef(
            metric_id="manual_result_quality",
            metric_name="人工评分：结果质量（0-5，可审计）",
            workbook_relpath="analysis_viz/data/derived/metrics/manual/manual_result_quality_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="manual_reasoning_quality",
            metric_name="人工评分：推理合理性（0-5，可审计）",
            workbook_relpath="analysis_viz/data/derived/metrics/manual/manual_reasoning_quality_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="alignment_result_vs_judge_case_level",
            metric_name="Alignment：人工结果质量 vs Judge(×5) 病例级",
            workbook_relpath="analysis_viz/data/derived/metrics/alignment/alignment_result_vs_judge_case_level_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="alignment_reasoning_vs_llm_case_level",
            metric_name="Alignment：人工推理合理性 vs LLM推理质量(×5) 病例级",
            workbook_relpath="analysis_viz/data/derived/metrics/alignment/alignment_reasoning_vs_llm_case_level_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="alignment_doctor_consensus_result",
            metric_name="Alignment：双医生结果质量一致性（病例级）",
            workbook_relpath="analysis_viz/data/derived/metrics/alignment/alignment_doctor_consensus_result_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
        MetricDef(
            metric_id="alignment_doctor_consensus_reasoning",
            metric_name="Alignment：双医生推理合理性一致性（病例级）",
            workbook_relpath="analysis_viz/data/derived/metrics/alignment/alignment_doctor_consensus_reasoning_source.xlsx",
            detail_sheets=("detail_used",),
            summary_sheets=("summary_used",),
            d4_stage_values=("D4_Plan",),
        ),
    ]

    rows: list[dict[str, str]] = []
    for d in metric_defs:
        workbook_path = ROOT / d.workbook_relpath
        if not workbook_path.exists():
            rows.append(
                {
                    "metric_id": d.metric_id,
                    "metric_name": d.metric_name,
                    "detail_relpath": "",
                    "summary_relpath": "",
                    "raw_source_relpath": "",
                    "source_sheet": "",
                    "includes_d1": "n/a",
                    "includes_gate3_d4": "n/a",
                    "n_cases": "0",
                    "n_excluded_d1": "0",
                    "n_excluded_gate3_d4": "0",
                    "notes": "missing workbook",
                }
            )
            continue

        wb = openpyxl.load_workbook(workbook_path, data_only=True)
        meta = _read_meta_kv(wb)

        raw_source = meta.get("raw_file") or meta.get("raw_source") or ""
        raw_sheet = meta.get("raw_sheet") or ""
        raw_source_rel = ""
        if raw_source:
            raw_source_rel = _rel(Path(raw_source))

        detail_refs = "|".join(
            f"{d.workbook_relpath}#{s}" for s in d.detail_sheets if s in wb.sheetnames
        )
        summary_refs = "|".join(
            f"{d.workbook_relpath}#{s}" for s in d.summary_sheets if s in wb.sheetnames
        )

        used_case_ids: set[str] = set()
        for sheet in d.detail_sheets:
            used_case_ids |= _sheet_case_ids(wb, sheet)

        excluded_d1 = _sheet_case_ids(wb, "excluded_d1_anomaly")
        excluded_gate3 = set()
        if "excluded_gate3_d4" in wb.sheetnames:
            excluded_gate3 = _sheet_case_ids(wb, "excluded_gate3_d4")
        if d.gate3_excluded_sheet and d.gate3_excluded_sheet in wb.sheetnames:
            excluded_gate3 = _sheet_case_ids(wb, d.gate3_excluded_sheet)

        # includes_*: True 表示“存在违规样本仍在 used 明细中”。
        includes_d1_str = "True" if _sheet_has_any_true(wb, d.detail_sheets[0], "is_d1_anomaly") else "False"

        violates_gate3_d4: bool | None = None
        if d.d4_stage_values:
            violates_gate3_d4 = _gate3_d4_violation_by_stage(wb, d.detail_sheets, d.d4_stage_values)
        if d.gate3_excluded_sheet and d.d4_column_name:
            violates_gate3_d4 = _gate3_d4_violation_by_blank_column(
                wb,
                detail_sheet=d.detail_sheets[0],
                excluded_sheet=d.gate3_excluded_sheet,
                d4_column_name=d.d4_column_name,
            )

        includes_gate3_d4_str = "n/a" if violates_gate3_d4 is None else ("True" if violates_gate3_d4 else "False")

        rows.append(
            {
                "metric_id": d.metric_id,
                "metric_name": d.metric_name,
                "detail_relpath": detail_refs,
                "summary_relpath": summary_refs,
                "raw_source_relpath": raw_source_rel,
                "source_sheet": raw_sheet,
                "includes_d1": includes_d1_str,
                "includes_gate3_d4": includes_gate3_d4_str,
                "n_cases": str(len(used_case_ids)),
                "n_excluded_d1": str(len(excluded_d1)),
                "n_excluded_gate3_d4": str(len(excluded_gate3)),
                "notes": "",
            }
        )

    METRICS_REGISTRY.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "metric_id",
        "metric_name",
        "detail_relpath",
        "summary_relpath",
        "raw_source_relpath",
        "source_sheet",
        "includes_d1",
        "includes_gate3_d4",
        "n_cases",
        "n_excluded_d1",
        "n_excluded_gate3_d4",
        "notes",
    ]
    with METRICS_REGISTRY.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def main() -> None:
    build_metrics_registry()


if __name__ == "__main__":
    main()
