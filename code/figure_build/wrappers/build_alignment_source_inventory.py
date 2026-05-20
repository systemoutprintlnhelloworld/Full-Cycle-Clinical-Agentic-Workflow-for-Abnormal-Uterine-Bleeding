"""Build alignment source-data observability inventory.

Purpose
-------
Classify each workbook under
`analysis_viz/figures/variants/outputs_latest__alignment/source_data`
by whether it has case-level detail columns (case_id/CaseID/病例ID),
and export a persistent inventory CSV for documentation/audit.

Output
------
- analysis_viz/docs/alignment_source_inventory.csv
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
ALIGN_SOURCE_DIR = ROOT / "analysis_viz" / "figures" / "variants" / "outputs_latest__alignment" / "source_data"
OUT_CSV = ROOT / "analysis_viz" / "docs" / "alignment_source_inventory.csv"

CASE_COLS = {"case_id", "CaseID", "病例ID"}
STAGE_COLS = {"stage", "Stage", "环节"}


def _detect_group(filename: str) -> str:
    lower = filename.lower()
    if lower.startswith("alignment_reasoning"):
        return "alignment_reasoning"
    if lower.startswith("alignment_result"):
        return "alignment_result"
    if lower.startswith("manual_reasoning"):
        return "manual_reasoning"
    if lower.startswith("manual_result"):
        return "manual_result"
    if lower.startswith("manual_top_score"):
        return "manual_top_score"
    return "other"


def _sheet_header(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[str]:
    if ws.max_row < 1:
        return []
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    out: list[str] = []
    for v in row:
        if v is None:
            continue
        out.append(str(v).strip())
    return out


def _sheet_data_rows(ws: openpyxl.worksheet.worksheet.Worksheet) -> int:
    return max(0, ws.max_row - 1)


def build_inventory() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    for path in sorted(ALIGN_SOURCE_DIR.glob("*.xlsx")):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            sheet_names = wb.sheetnames
            has_case_level = False
            has_stage_col = False
            header_union: set[str] = set()
            sheet_rows: list[str] = []

            for name in sheet_names:
                ws = wb[name]
                header = _sheet_header(ws)
                header_union.update(header)
                if any(c in CASE_COLS for c in header):
                    has_case_level = True
                if any(c in STAGE_COLS for c in header):
                    has_stage_col = True
                sheet_rows.append(f"{name}:{_sheet_data_rows(ws)}")

            if has_case_level:
                observability = "case_level"
            elif has_stage_col:
                observability = "stage_or_model_summary_only"
            else:
                observability = "metadata_or_aggregate_only"

            rows.append(
                {
                    "file_relpath": str(path.relative_to(ROOT)).replace("\\", "/"),
                    "filename": path.name,
                    "group": _detect_group(path.name),
                    "observability_level": observability,
                    "has_case_level": str(has_case_level),
                    "has_stage_column": str(has_stage_col),
                    "sheet_names": "|".join(sheet_names),
                    "sheet_rows": "|".join(sheet_rows),
                    "header_union": "|".join(sorted(header_union)),
                    "notes": "",
                }
            )
        finally:
            wb.close()

    return rows


def write_csv(rows: list[dict[str, str]]) -> None:
    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "file_relpath",
        "filename",
        "group",
        "observability_level",
        "has_case_level",
        "has_stage_column",
        "sheet_names",
        "sheet_rows",
        "header_union",
        "notes",
    ]
    with OUT_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def main() -> None:
    rows = build_inventory()
    write_csv(rows)
    print("WROTE", OUT_CSV)
    print("ROWS", len(rows))


if __name__ == "__main__":
    main()

