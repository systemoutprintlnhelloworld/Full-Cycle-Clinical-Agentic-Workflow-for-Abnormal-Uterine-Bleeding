"""Audit observability schema of derived machine-metric workbooks."""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
METRIC_DIR = ROOT / "analysis_viz" / "data" / "derived" / "metrics"
OUT_CSV = ROOT / "analysis_viz" / "docs" / "metric_observability_schema_audit.csv"
OUT_MD = ROOT / "analysis_viz" / "docs" / "metric_observability_schema_audit.md"


def _header_set(ws: openpyxl.worksheet.worksheet.Worksheet) -> set[str]:
    if ws.max_row < 1:
        return set()
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    out: set[str] = set()
    for v in row:
        if v is None:
            continue
        out.add(str(v).strip())
    return out


def _sheet_rows(ws: openpyxl.worksheet.worksheet.Worksheet) -> int:
    return max(0, ws.max_row - 1)


def main() -> None:
    rows: list[dict[str, str]] = []

    for path in sorted(METRIC_DIR.rglob("*.xlsx")):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            names = wb.sheetnames
            detail_sheet_candidates = [s for s in ["detail_used", "fact_detail_used", "cross_detail_used"] if s in names]
            has_detail = len(detail_sheet_candidates) > 0
            has_summary = "summary_used" in names
            has_ex_d1 = "excluded_d1_anomaly" in names
            has_ex_gate = "excluded_gate3_d4" in names
            has_meta = "meta" in names

            detail_rows = 0
            summary_rows = 0
            has_model_short = False
            has_stage_cn = False
            has_source_file = False
            has_stage_col = False
            has_d4_hint = False

            for s in detail_sheet_candidates:
                ws = wb[s]
                detail_rows += _sheet_rows(ws)
                hdr = _header_set(ws)
                has_model_short = has_model_short or ("model_short" in hdr)
                has_stage_cn = has_stage_cn or ("stage_cn" in hdr)
                has_source_file = has_source_file or ("source_file" in hdr)
                has_stage_col = has_stage_col or ("stage" in hdr)
                has_d4_hint = has_d4_hint or any("D4" in h for h in hdr)

            if has_summary:
                summary_rows = _sheet_rows(wb["summary_used"])

            gate3_required = has_d4_hint or has_stage_col
            gate3_ok = has_ex_gate if gate3_required else True
            stage_cn_ok = has_stage_cn if has_stage_col else True

            score = 0
            for b in [has_detail, has_summary, has_ex_d1, gate3_ok, has_meta, has_model_short, stage_cn_ok]:
                score += 1 if b else 0

            rows.append(
                {
                    "workbook_relpath": str(path.relative_to(ROOT)).replace("\\", "/"),
                    "has_detail_used": str(has_detail),
                    "has_summary_used": str(has_summary),
                    "has_excluded_d1": str(has_ex_d1),
                    "has_excluded_gate3_d4": str(has_ex_gate),
                    "gate3_required": str(gate3_required),
                    "has_meta": str(has_meta),
                    "detail_rows": str(detail_rows),
                    "summary_rows": str(summary_rows),
                    "has_model_short": str(has_model_short),
                    "has_stage_cn": str(has_stage_cn),
                    "has_source_file": str(has_source_file),
                    "schema_score_0_7": str(score),
                    "notes": "",
                }
            )
        finally:
            wb.close()

    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_CSV.open("w", encoding="utf-8", newline="") as f:
        fields = [
            "workbook_relpath",
            "has_detail_used",
            "has_summary_used",
            "has_excluded_d1",
            "has_excluded_gate3_d4",
            "gate3_required",
            "has_meta",
            "detail_rows",
            "summary_rows",
            "has_model_short",
            "has_stage_cn",
            "has_source_file",
            "schema_score_0_7",
            "notes",
        ]
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow(r)

    total = len(rows)
    high = sum(1 for r in rows if int(r["schema_score_0_7"]) >= 6)
    low_items = [r for r in rows if int(r["schema_score_0_7"]) < 6]

    lines = [
        "# 机器指标可观测性结构审计",
        "",
        f"- 工作簿数量：{total}",
        f"- 高可观测（score>=6）数量：{high}",
        f"- 需改进（score<6）数量：{len(low_items)}",
        "",
        "## 需改进项",
        "",
    ]
    if low_items:
        for r in low_items:
            lines.append(f"- `{r['workbook_relpath']}` (score={r['schema_score_0_7']})")
    else:
        lines.append("- 无")

    lines += ["", "## 详表", "", f"- `analysis_viz/docs/{OUT_CSV.name}`", ""]
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")

    print("WROTE", OUT_CSV)
    print("WROTE", OUT_MD)
    print("TOTAL", total)
    print("LOW", len(low_items))


if __name__ == "__main__":
    main()

