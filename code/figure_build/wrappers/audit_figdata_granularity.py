"""Audit figdata granularity (case-level vs summary-level).

Outputs:
- analysis_viz/docs/figdata_granularity_audit.csv
- analysis_viz/docs/figdata_granularity_audit.md
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[3]

PAPER_FIGDATA_DIR = ROOT / "论文" / "figdata"
DERIVED_PAPER_FIGDATA_DIR = ROOT / "analysis_viz" / "data" / "derived" / "figdata" / "paper"

OUT_CSV = ROOT / "analysis_viz" / "docs" / "figdata_granularity_audit.csv"
OUT_MD = ROOT / "analysis_viz" / "docs" / "figdata_granularity_audit.md"

CASE_COLS = {"case_id", "CaseID", "病例ID"}
STAGE_COLS = {"stage", "Stage", "阶段", "环节"}
CENTER_COLS = {"center", "Center", "中心"}
MODEL_COLS = {"model", "Model", "模型", "模型名称", "被评测模型"}


def _prefix_from_name(name: str) -> str:
    if "__" in name:
        return name.split("__", 1)[0]
    if "_" in name and name.lower().startswith("fig"):
        return name.split("_", 1)[0]
    if name.lower().startswith("fig"):
        return name.split(".", 1)[0]
    return "UNKNOWN"


def _classify_level(*, has_case: bool, has_stage: bool, has_center: bool, has_model: bool, file_ext: str) -> str:
    if has_case:
        return "case_level"
    if has_stage or has_center or has_model:
        return "stage_or_model_summary_only"
    if file_ext == ".json":
        return "json_or_metadata_only"
    return "aggregate_only"


def _audit_csv(path: Path) -> dict[str, object]:
    try:
        frame = pd.read_csv(path, nrows=5)
    except Exception:
        return {
            "sheet_count": 1,
            "total_rows": -1,
            "header_union": "",
            "has_case_id": False,
            "has_stage": False,
            "has_center": False,
            "has_model": False,
            "notes": "csv_read_failed",
        }

    columns = [str(column).strip() for column in frame.columns]
    column_set = set(columns)
    return {
        "sheet_count": 1,
        "total_rows": _safe_count_rows_csv(path),
        "header_union": "|".join(columns[:40]),
        "has_case_id": bool(column_set & CASE_COLS),
        "has_stage": bool(column_set & STAGE_COLS),
        "has_center": bool(column_set & CENTER_COLS),
        "has_model": bool(column_set & MODEL_COLS),
        "notes": "",
    }


def _safe_count_rows_csv(path: Path) -> int:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as file_handle:
            return max(0, sum(1 for _ in file_handle) - 1)
    except Exception:
        return -1


def _audit_xlsx(path: Path) -> dict[str, object]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        headers: set[str] = set()
        total_rows = 0
        has_case = False
        has_stage = False
        has_center = False
        has_model = False

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            total_rows += max(0, worksheet.max_row - 1)
            if worksheet.max_row < 1:
                continue

            first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
            row_headers = {str(value).strip() for value in first_row if value is not None and str(value).strip()}
            headers |= row_headers
            has_case = has_case or bool(row_headers & CASE_COLS)
            has_stage = has_stage or bool(row_headers & STAGE_COLS)
            has_center = has_center or bool(row_headers & CENTER_COLS)
            has_model = has_model or bool(row_headers & MODEL_COLS)

        return {
            "sheet_count": len(workbook.sheetnames),
            "total_rows": total_rows,
            "header_union": "|".join(sorted(headers)[:60]),
            "has_case_id": has_case,
            "has_stage": has_stage,
            "has_center": has_center,
            "has_model": has_model,
            "notes": "",
        }
    finally:
        workbook.close()


def _audit_json(path: Path) -> dict[str, object]:
    return {
        "sheet_count": 1,
        "total_rows": -1,
        "header_union": "",
        "has_case_id": False,
        "has_stage": False,
        "has_center": False,
        "has_model": False,
        "notes": "json_not_columnar",
    }


def _iter_figdata_files() -> list[tuple[str, Path]]:
    files: list[tuple[str, Path]] = []
    for scope, base_dir in [
        ("paper_figdata", PAPER_FIGDATA_DIR),
        ("derived_paper_figdata", DERIVED_PAPER_FIGDATA_DIR),
    ]:
        if not base_dir.exists():
            continue
        for path in sorted(base_dir.glob("*")):
            if not path.is_file():
                continue
            if path.suffix.lower() not in {".csv", ".xlsx", ".json"}:
                continue
            files.append((scope, path))
    return files


def main() -> None:
    rows: list[dict[str, object]] = []

    for scope, file_path in _iter_figdata_files():
        extension = file_path.suffix.lower()
        if extension == ".csv":
            info = _audit_csv(file_path)
        elif extension == ".xlsx":
            info = _audit_xlsx(file_path)
        else:
            info = _audit_json(file_path)

        observability_level = _classify_level(
            has_case=bool(info["has_case_id"]),
            has_stage=bool(info["has_stage"]),
            has_center=bool(info["has_center"]),
            has_model=bool(info["has_model"]),
            file_ext=extension,
        )
        rows.append(
            {
                "scope": scope,
                "file_relpath": str(file_path.relative_to(ROOT)).replace("\\", "/"),
                "filename": file_path.name,
                "figure_prefix": _prefix_from_name(file_path.name),
                "file_ext": extension,
                "sheet_count": info["sheet_count"],
                "total_rows": info["total_rows"],
                "has_case_id": info["has_case_id"],
                "has_stage": info["has_stage"],
                "has_center": info["has_center"],
                "has_model": info["has_model"],
                "observability_level": observability_level,
                "header_union": info["header_union"],
                "notes": info["notes"],
            }
        )

    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "scope",
        "file_relpath",
        "filename",
        "figure_prefix",
        "file_ext",
        "sheet_count",
        "total_rows",
        "has_case_id",
        "has_stage",
        "has_center",
        "has_model",
        "observability_level",
        "header_union",
        "notes",
    ]
    with OUT_CSV.open("w", encoding="utf-8", newline="") as file_handle:
        writer = csv.DictWriter(file_handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)

    frame = pd.DataFrame(rows)
    by_level = frame["observability_level"].value_counts().to_dict() if not frame.empty else {}
    by_prefix = (
        frame.groupby(["figure_prefix", "observability_level"]).size().reset_index(name="count")
        if not frame.empty
        else pd.DataFrame(columns=["figure_prefix", "observability_level", "count"])
    )

    markdown_lines = [
        "# figdata 粒度审计",
        "",
        f"- 审计文件数：{len(rows)}",
        "",
        "## 按粒度分布",
        "",
    ]
    if by_level:
        for level, count in sorted(by_level.items(), key=lambda item: item[0]):
            markdown_lines.append(f"- `{level}`: {count}")
    else:
        markdown_lines.append("- 无")

    markdown_lines += ["", "## 按图前缀分布（prefix + level）", ""]
    if not by_prefix.empty:
        for _, row in by_prefix.sort_values(["figure_prefix", "observability_level"]).iterrows():
            markdown_lines.append(
                f"- `{row['figure_prefix']}` / `{row['observability_level']}`: {int(row['count'])}"
            )
    else:
        markdown_lines.append("- 无")

    markdown_lines += ["", "## 详表", "", f"- `analysis_viz/docs/{OUT_CSV.name}`", ""]
    OUT_MD.write_text("\n".join(markdown_lines), encoding="utf-8")

    print("WROTE", OUT_CSV)
    print("WROTE", OUT_MD)
    print("ROWS", len(rows))


if __name__ == "__main__":
    main()

