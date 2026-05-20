from __future__ import annotations

import json
import hashlib
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd

from .checks import build_check_rounds
from .discovery import discover_datasets
from .excel_export import ExcelSheetSpec, export_metrics_source_data_xlsx
from .field_audit import (
    collect_column_stats,
    compute_common_columns,
    scan_column_anomalies,
    scan_missing_value_anomalies,
)
from .fields import (
    build_d1_decision_fields,
    build_d2_decision_fields,
    build_d3_decision_fields,
    build_d4_rehab_fields,
    build_judge_scores_by_case,
)
from .ingest import load_dataset
from .llm_tasks import (
    generate_d1_diagnosis_topk_tasks,
    generate_diagnosis_quality_tasks,
    generate_diagnosis_bias_tasks,
    generate_cross_stage_consistency_tasks,
    generate_final_diagnosis_proximity_tasks,
    generate_gate1_dx_rescore_tasks,
    generate_gate2_dx_plan_rescore_tasks,
    generate_fact_consistency_and_missing_tasks,
    generate_memory_retention_tasks,
    generate_plan_quality_tasks,
    generate_rehab_followup_quality_tasks,
    generate_rationale_quality_tasks,
    generate_unmatched_check_reasonableness_tasks,
    write_tasks_jsonl,
)
from .metrics import MetricTables, build_metric_tables
from .overrides import (
    apply_check_round_overrides,
    apply_review_feedback_to_override_template,
    load_check_round_overrides_csv,
    update_check_round_override_template,
)
from .quality import canonicalize_sheet_columns, scan_sheet_anomalies
from .repairs import repair_doc_decision_nested_fields, repair_doc_loop_decision_duplicate_raw_json
from .review_export import build_review_workbook_spec, export_review_workbook_xlsx
from .sorting import case_id_sort_key
from .status_audit import run_status_audit
from .types import LoadedData

# LLM tasks currently confirmed in metrics-overview (others are skipped to reduce cost).
LLM_TASK_ALLOWLIST: set[str] = {
    "llm.rationale_quality",
    "llm.fact_consistency_and_missing",
    "llm.memory_retention",
    "llm.cross_stage_consistency",
    "llm.final_diagnosis_proximity",
    "llm.gate1_dx_rescore",
    "llm.gate2_dx_plan_rescore",
}


def _make_run_id(tag: str) -> str:
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    safe_tag = "".join(ch if ch.isalnum() or ch in {"-", "_"} else "-" for ch in tag.strip().lower()).strip("-")
    safe_tag = safe_tag or "metrics"
    if safe_tag in {"latest", "current"}:
        return "latest"
    return f"{ts}_{safe_tag}"


def run_metrics(
    project_root: Path,
    data_root: Path,
    outputs_root: Path,
    work_root: Path,
    tag: str,
    centers: list[str] | None,
    models: list[str] | None,
    match_score_threshold: float,
    export_excel: bool,
    split_by: str = "center-model",
    llm_task_scope: str = "review_only",
) -> None:
    data_root = (project_root / data_root).resolve()
    outputs_root = (project_root / outputs_root).resolve()
    work_root = (project_root / work_root).resolve()

    run_id = _make_run_id(tag)
    out_dir = outputs_root / run_id
    work_dir = work_root / run_id
    if run_id == "latest":
        # User requirement: keep history (do NOT delete), but still avoid mixing old/new.
        # Strategy: move previous generated artifacts into `outputs/archive/<ts>_latest/` and `work/archive/<ts>_latest/`.

        def _safe_move_tree(src: Path, dst: Path) -> None:
            """
            Move files/dirs while tolerating Excel lock files (~$).
            If a file is locked, skip it and continue.
            """
            if not src.exists():
                return
            dst.mkdir(parents=True, exist_ok=True)
            for child in list(src.iterdir()):
                # Skip Excel temp/lock files
                if child.is_file() and child.name.startswith("~$"):
                    continue
                target = dst / child.name
                try:
                    shutil.move(str(child), str(target))
                except PermissionError:
                    # Leave locked files in place; continue.
                    continue
                except Exception:
                    continue

        def _unique_dest(base: Path) -> Path:
            if not base.exists():
                return base
            for i in range(1, 1000):
                cand = base.parent / f"{base.name}-{i}"
                if not cand.exists():
                    return cand
            raise RuntimeError(f"Cannot find unique archive dest for {base}")

        ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

        # Archive outputs/latest (but keep summary/llm_results_*.xlsx in-place for cost control).
        if out_dir.exists():
            out_archive_root = outputs_root / "archive"
            out_archive_root.mkdir(parents=True, exist_ok=True)
            out_archive_dir = _unique_dest(out_archive_root / f"{ts}_latest")
            out_archive_dir.mkdir(parents=True, exist_ok=True)

            summary_keep: set[str] = set()
            summary_dir = out_dir / "summary"
            if summary_dir.exists():
                for f in summary_dir.glob("llm_results_*.xlsx"):
                    summary_keep.add(f.name)

            for child in list(out_dir.iterdir()):
                if child.name == "summary" and child.is_dir():
                    # Move summary children except cached llm_results
                    dst_summary = out_archive_dir / "summary"
                    dst_summary.mkdir(parents=True, exist_ok=True)
                    for s_child in list(child.iterdir()):
                        if s_child.is_file() and s_child.name in summary_keep:
                            continue
                        if s_child.is_file() and s_child.name.startswith("~$"):
                            continue
                        try:
                            shutil.move(str(s_child), str(dst_summary / s_child.name))
                        except PermissionError:
                            continue
                        except Exception:
                            continue
                    continue
                if child.is_dir():
                    _safe_move_tree(child, out_archive_dir / child.name)
                else:
                    if child.name.startswith("~$"):
                        continue
                    try:
                        shutil.move(str(child), str(out_archive_dir / child.name))
                    except PermissionError:
                        continue
                    except Exception:
                        continue

        # Archive work/latest (keep user-editable + LLM cache).
        if work_dir.exists():
            work_archive_root = work_root / "archive"
            work_archive_root.mkdir(parents=True, exist_ok=True)
            work_archive_dir = _unique_dest(work_archive_root / f"{ts}_latest")
            work_archive_dir.mkdir(parents=True, exist_ok=True)

            keep_names = {"llm_results", "overrides", "feedback", "repairs", "任务追踪.csv"}
            for child in list(work_dir.iterdir()):
                if child.name in keep_names:
                    continue
                if child.is_dir():
                    _safe_move_tree(child, work_archive_dir / child.name)
                else:
                    if child.name.startswith("~$"):
                        continue
                    try:
                        shutil.move(str(child), str(work_archive_dir / child.name))
                    except PermissionError:
                        continue
                    except Exception:
                        continue

        out_dir.mkdir(parents=True, exist_ok=True)
        work_dir.mkdir(parents=True, exist_ok=True)
    else:
        out_dir.mkdir(parents=True, exist_ok=True)
        work_dir.mkdir(parents=True, exist_ok=True)

    discovery = discover_datasets(data_root, centers=centers, models=models)
    (work_dir / "discovery_warnings.json").write_text(json.dumps(discovery.warnings, ensure_ascii=False, indent=2), encoding="utf-8")

    if split_by == "center-model":
        _run_metrics_split_by_center_model(
            project_root=project_root,
            out_dir=out_dir,
            work_dir=work_dir,
            discovery=discovery,
            match_score_threshold=match_score_threshold,
            export_excel=export_excel,
            llm_task_scope=llm_task_scope,
        )
        # Keep "latest" self-contained for human review: also generate fixed excel copies with inferred cross-sheet status.
        if run_id == "latest":
            run_status_audit(
                project_root=project_root,
                data_root=Path("data"),
                work_root=Path("work"),
                tag="latest",
                centers=centers,
                models=models,
            )
        if out_dir.name != "latest":
            _update_latest_outputs_index(outputs_root=outputs_root, out_dir=out_dir)
        print(str(out_dir))
        return

    if split_by not in {"none", "center-model"}:
        raise ValueError(f"Unknown split_by={split_by!r}. Expected 'none' or 'center-model'.")

    all_case_index: list[pd.DataFrame] = []
    all_check_rounds: list[pd.DataFrame] = []
    all_gt_checks: list[pd.DataFrame] = []
    all_metrics_by_case: list[pd.DataFrame] = []
    all_metrics_by_center_model: list[pd.DataFrame] = []
    all_llm_tasks_dx_topk_d1: list[dict] = []
    all_llm_tasks_rationale: list[dict] = []
    all_llm_tasks_fact_consistency: list[dict] = []
    all_llm_tasks_memory_retention: list[dict] = []
    all_llm_tasks_cross_stage: list[dict] = []
    all_llm_tasks_final_dx_proximity: list[dict] = []
    all_llm_tasks_gate1_rescore: list[dict] = []
    all_llm_tasks_gate2_rescore: list[dict] = []
    all_d1_decision_fields: list[pd.DataFrame] = []
    all_d2_decision_fields: list[pd.DataFrame] = []
    all_d3_decision_fields: list[pd.DataFrame] = []
    all_d4_rehab_fields: list[pd.DataFrame] = []
    all_judge_scores: list[pd.DataFrame] = []

    for fileset in discovery.datasets:
        loaded = load_dataset(fileset)
        check_rounds = build_check_rounds(loaded, match_score_threshold=match_score_threshold)
        tables: MetricTables = build_metric_tables(loaded, check_rounds=check_rounds)

        all_case_index.append(tables.case_index)
        all_check_rounds.append(tables.check_rounds)
        all_gt_checks.append(tables.gt_checks)
        all_metrics_by_case.append(tables.metrics_by_case)
        all_metrics_by_center_model.append(tables.metrics_by_center_model)

        # LLM task generation (no API call here; just prepare tasks JSONL).
        if "llm.diagnosis_semantic_match" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_dx_topk_d1.extend(generate_d1_diagnosis_topk_tasks(project_root, loaded))
        if "llm.rationale_quality" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_rationale.extend(generate_rationale_quality_tasks(project_root, loaded))
        if "llm.fact_consistency_and_missing" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_fact_consistency.extend(generate_fact_consistency_and_missing_tasks(project_root, loaded))
        if "llm.memory_retention" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_memory_retention.extend(generate_memory_retention_tasks(project_root, loaded))
        if "llm.cross_stage_consistency" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_cross_stage.extend(generate_cross_stage_consistency_tasks(project_root, loaded))
        if "llm.final_diagnosis_proximity" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_final_dx_proximity.extend(generate_final_diagnosis_proximity_tasks(project_root, loaded))
        if "llm.gate1_dx_rescore" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_gate1_rescore.extend(generate_gate1_dx_rescore_tasks(project_root, loaded))
        if "llm.gate2_dx_plan_rescore" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_gate2_rescore.extend(generate_gate2_dx_plan_rescore_tasks(project_root, loaded))

        # Field packs (for human review / Excel analysis)
        all_d1_decision_fields.append(build_d1_decision_fields(loaded))
        all_d2_decision_fields.append(build_d2_decision_fields(loaded))
        all_d3_decision_fields.append(build_d3_decision_fields(loaded))
        all_d4_rehab_fields.append(build_d4_rehab_fields(loaded))
        all_judge_scores.append(build_judge_scores_by_case(loaded))

    case_index = pd.concat(all_case_index, ignore_index=True) if all_case_index else pd.DataFrame()
    check_rounds = pd.concat(all_check_rounds, ignore_index=True) if all_check_rounds else pd.DataFrame()
    gt_checks = pd.concat(all_gt_checks, ignore_index=True) if all_gt_checks else pd.DataFrame()
    metrics_by_case = pd.concat(all_metrics_by_case, ignore_index=True) if all_metrics_by_case else pd.DataFrame()
    metrics_by_center_model = pd.concat(all_metrics_by_center_model, ignore_index=True) if all_metrics_by_center_model else pd.DataFrame()
    d1_decision_fields = pd.concat(all_d1_decision_fields, ignore_index=True) if all_d1_decision_fields else pd.DataFrame()
    d2_decision_fields = pd.concat(all_d2_decision_fields, ignore_index=True) if all_d2_decision_fields else pd.DataFrame()
    d3_decision_fields = pd.concat(all_d3_decision_fields, ignore_index=True) if all_d3_decision_fields else pd.DataFrame()
    d4_rehab_fields = pd.concat(all_d4_rehab_fields, ignore_index=True) if all_d4_rehab_fields else pd.DataFrame()
    judge_scores = pd.concat(all_judge_scores, ignore_index=True) if all_judge_scores else pd.DataFrame()

    check_rounds_audit_summary = pd.DataFrame()
    check_rounds_review = pd.DataFrame()
    d1_decision_anomalies = pd.DataFrame()

    # Enrich metrics_by_case with compact judge scores (numeric-friendly).
    if not metrics_by_case.empty and not judge_scores.empty:
        metrics_by_case = metrics_by_case.merge(judge_scores, on=["center", "model", "case_id"], how="left")

    # Write CSV outputs (machine-friendly)
    case_index.to_csv(out_dir / "case_index.csv", index=False, encoding="utf-8-sig")
    check_rounds.to_csv(out_dir / "check_rounds.csv", index=False, encoding="utf-8-sig")
    gt_checks.to_csv(out_dir / "gt_checks.csv", index=False, encoding="utf-8-sig")
    metrics_by_case.to_csv(out_dir / "metrics_by_case.csv", index=False, encoding="utf-8-sig")
    metrics_by_center_model.to_csv(out_dir / "metrics_by_center_model.csv", index=False, encoding="utf-8-sig")
    if not d1_decision_fields.empty:
        d1_decision_fields.to_csv(out_dir / "d1_decision_fields.csv", index=False, encoding="utf-8-sig")
    if not d2_decision_fields.empty:
        d2_decision_fields.to_csv(out_dir / "d2_decision_fields.csv", index=False, encoding="utf-8-sig")
    if not d3_decision_fields.empty:
        d3_decision_fields.to_csv(out_dir / "d3_decision_fields.csv", index=False, encoding="utf-8-sig")
    if not d4_rehab_fields.empty:
        d4_rehab_fields.to_csv(out_dir / "d4_rehab_fields.csv", index=False, encoding="utf-8-sig")
    if not judge_scores.empty:
        judge_scores.to_csv(out_dir / "judge_scores_by_case.csv", index=False, encoding="utf-8-sig")

    # Audit helpers (human-in-the-loop): summarize suspicious check-round rows for review.
    if not check_rounds.empty:
        tmp = check_rounds.copy()
        # Evidence-driven: a round is executed if AI requested checks in that round.
        tmp["executed_round"] = pd.to_numeric(tmp["ai_total_requested_count"], errors="coerce").fillna(0) > 0
        tmp["has_extra_suspect"] = tmp["judge_extra_matched_items_suspect"].astype(str).str.strip().ne("[]")
        tmp["judge_count_vs_unexec_conflict"] = tmp["judge_count_vs_unexec_conflict"].fillna(False)
        tmp["ai_total_vs_judge_total_mismatch"] = tmp["ai_total_vs_judge_total_mismatch"].fillna(False)

        check_rounds_audit_summary = tmp.groupby(["center", "model", "stage"], as_index=False).agg(
            executed_round_rows=("executed_round", "sum"),
            conflict_rows=("judge_count_vs_unexec_conflict", "sum"),
            total_mismatch_rows=("ai_total_vs_judge_total_mismatch", "sum"),
            extra_suspect_rows=("has_extra_suspect", "sum"),
        )
        check_rounds_audit_summary.to_csv(out_dir / "check_rounds_audit_summary.csv", index=False, encoding="utf-8-sig")

        review_mask = tmp["executed_round"] & (
            tmp["judge_count_vs_unexec_conflict"] | tmp["ai_total_vs_judge_total_mismatch"] | tmp["has_extra_suspect"]
        )
        review_cols = [
            "center",
            "model",
            "case_id",
            "stage",
            "round_idx",
            "doc_status",
            "judge_status",
            "doc_request_col",
            "doc_ai_requests_text",
            "ai_total_requested_count",
            "matched_count_final",
            "matched_count_source",
            "unmatched_count_from_unexecuted_list",
            "unmatched_count_final",
            "judge_matched_count_raw",
            "judge_total_requested_count_raw",
            "judge_match_score_raw",
            "judge_reasonable_score_raw",
            "judge_matched_items_raw",
            "judge_unexecuted_items_raw",
            "judge_extra_matched_items_suspect",
            "judge_count_vs_unexec_conflict",
            "ai_total_vs_judge_total_mismatch",
        ]
        check_rounds_review = tmp.loc[review_mask, [c for c in review_cols if c in tmp.columns]].copy()
        check_rounds_review.to_csv(out_dir / "check_rounds_review.csv", index=False, encoding="utf-8-sig")

    # D1 decision anomalies list (for re-judge / special handling)
    if not case_index.empty and "is_anomaly_d1_to_d2" in case_index.columns:
        d1_decision_anomalies = case_index.loc[case_index["is_anomaly_d1_to_d2"] == True, ["center", "model", "case_id"]].copy()
        d1_decision_anomalies.to_csv(out_dir / "d1_decision_anomalies.csv", index=False, encoding="utf-8-sig")

    if all_llm_tasks_dx_topk_d1:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.diagnosis_semantic_match.D1.jsonl", all_llm_tasks_dx_topk_d1)
    if all_llm_tasks_rationale:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.rationale_quality.jsonl", all_llm_tasks_rationale)
    if all_llm_tasks_fact_consistency:
        write_tasks_jsonl(
            work_dir / "llm_tasks" / "llm.fact_consistency_and_missing.jsonl",
            all_llm_tasks_fact_consistency,
        )
    if all_llm_tasks_memory_retention:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.memory_retention.jsonl", all_llm_tasks_memory_retention)
    if all_llm_tasks_cross_stage:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.cross_stage_consistency.jsonl", all_llm_tasks_cross_stage)
    if all_llm_tasks_final_dx_proximity:
        write_tasks_jsonl(
            work_dir / "llm_tasks" / "llm.final_diagnosis_proximity.jsonl",
            all_llm_tasks_final_dx_proximity,
        )
    if all_llm_tasks_gate1_rescore:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.gate1_dx_rescore.jsonl", all_llm_tasks_gate1_rescore)
    if all_llm_tasks_gate2_rescore:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.gate2_dx_plan_rescore.jsonl", all_llm_tasks_gate2_rescore)

    if export_excel:
        export_metrics_source_data_xlsx(
            out_dir / "metrics_source_data.xlsx",
            sheets=[
                ExcelSheetSpec("case_index", case_index),
                ExcelSheetSpec("check_rounds", check_rounds),
                ExcelSheetSpec("check_rounds_audit_summary", check_rounds_audit_summary),
                ExcelSheetSpec("check_rounds_review", check_rounds_review),
                ExcelSheetSpec("d1_decision_anomalies", d1_decision_anomalies),
                ExcelSheetSpec("gt_checks", gt_checks),
                ExcelSheetSpec("judge_scores_by_case", judge_scores),
                ExcelSheetSpec("d1_decision_fields", d1_decision_fields),
                ExcelSheetSpec("d2_decision_fields", d2_decision_fields),
                ExcelSheetSpec("d3_decision_fields", d3_decision_fields),
                ExcelSheetSpec("d4_rehab_fields", d4_rehab_fields),
                ExcelSheetSpec("metrics_by_case", metrics_by_case),
                ExcelSheetSpec("metrics_by_center_model", metrics_by_center_model),
            ],
        )

    # Minimal console output (for CLI)
    print(str(out_dir))


def _build_check_rounds_audit_tables(check_rounds: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    if check_rounds.empty:
        return pd.DataFrame(), pd.DataFrame()

    tmp = check_rounds.copy()
    tmp["executed_round"] = pd.to_numeric(tmp["ai_total_requested_count"], errors="coerce").fillna(0) > 0
    tmp["missing_judge"] = tmp["executed_round"] & pd.to_numeric(tmp["matched_count_final"], errors="coerce").isna()
    tmp["has_extra_suspect"] = tmp["judge_extra_matched_items_suspect"].astype(str).str.strip().ne("[]")
    tmp["judge_count_vs_unexec_conflict"] = tmp["judge_count_vs_unexec_conflict"].fillna(False)
    tmp["ai_total_vs_judge_total_mismatch"] = tmp["ai_total_vs_judge_total_mismatch"].fillna(False)
    tmp["review_reasons"] = ""
    tmp.loc[tmp["missing_judge"] == True, "review_reasons"] += "裁判缺失/未评测；"
    tmp.loc[tmp["judge_count_vs_unexec_conflict"] == True, "review_reasons"] += "计数冲突；"
    tmp.loc[tmp["ai_total_vs_judge_total_mismatch"] == True, "review_reasons"] += "AI请求数与判官总数不一致；"
    tmp.loc[tmp["has_extra_suspect"] == True, "review_reasons"] += "疑似额外匹配项；"
    tmp["review_reasons"] = tmp["review_reasons"].astype(str).str.rstrip("；")

    check_rounds_audit_summary = tmp.groupby(["center", "model", "stage"], as_index=False).agg(
        executed_round_rows=("executed_round", "sum"),
        missing_judge_rows=("missing_judge", "sum"),
        conflict_rows=("judge_count_vs_unexec_conflict", "sum"),
        total_mismatch_rows=("ai_total_vs_judge_total_mismatch", "sum"),
        extra_suspect_rows=("has_extra_suspect", "sum"),
    )

    review_mask = tmp["executed_round"] & (
        tmp["missing_judge"] | tmp["judge_count_vs_unexec_conflict"] | tmp["ai_total_vs_judge_total_mismatch"] | tmp["has_extra_suspect"]
    )
    review_cols = [
        "center",
        "model",
        "case_id",
        "stage",
        "round_idx",
        "review_reasons",
        "doc_status",
        "judge_status",
        "doc_request_col",
        "doc_ai_requests_text",
        "ai_total_requested_count",
        "ai_total_source",
        "matched_count_final",
        "matched_count_source",
        "unmatched_count_from_unexecuted_list",
        "unmatched_count_final",
        "judge_matched_count_raw",
        "judge_total_requested_count_raw",
        "judge_match_score_raw",
        "judge_reasonable_score_raw",
        "judge_matched_items_raw",
        "judge_unexecuted_items_raw",
        "judge_extra_matched_items_suspect",
        "judge_count_vs_unexec_conflict",
        "ai_total_vs_judge_total_mismatch",
    ]
    check_rounds_review = tmp.loc[review_mask, [c for c in review_cols if c in tmp.columns]].copy()
    return check_rounds_audit_summary, check_rounds_review


def _run_metrics_split_by_center_model(
    project_root: Path,
    out_dir: Path,
    work_dir: Path,
    discovery,
    match_score_threshold: float,
    export_excel: bool,
    llm_task_scope: str,
) -> None:
    summary_dir = out_dir / "summary"
    summary_dir.mkdir(parents=True, exist_ok=True)

    # Pre-scan column stats to detect missing/extra columns (doc/judge) per sheet.
    column_stats, dataset_cols_map, column_totals = collect_column_stats(discovery.datasets)
    common_cols = compute_common_columns(column_stats, column_totals)

    metrics_by_center_model_rows: list[pd.DataFrame] = []
    metrics_by_case_rows: list[pd.DataFrame] = []
    metrics_by_case_full_rows: list[pd.DataFrame] = []
    case_index_rows: list[pd.DataFrame] = []
    check_rounds_rows: list[pd.DataFrame] = []
    check_rounds_audit_rows: list[pd.DataFrame] = []
    d1_decision_anomalies_rows: list[pd.DataFrame] = []
    gt_checks_rows: list[pd.DataFrame] = []
    judge_scores_rows: list[pd.DataFrame] = []
    d1_decision_fields_rows: list[pd.DataFrame] = []
    d2_decision_fields_rows: list[pd.DataFrame] = []
    d3_decision_fields_rows: list[pd.DataFrame] = []
    d4_rehab_fields_rows: list[pd.DataFrame] = []
    all_llm_tasks_dx_topk_d1: list[dict] = []
    all_llm_tasks_rationale: list[dict] = []
    all_llm_tasks_fact_consistency: list[dict] = []
    all_llm_tasks_memory_retention: list[dict] = []
    all_llm_tasks_cross_stage: list[dict] = []
    all_llm_tasks_unmatched_checks: list[dict] = []
    all_llm_tasks_dx_quality: list[dict] = []
    all_llm_tasks_plan_quality: list[dict] = []
    all_llm_tasks_dx_bias: list[dict] = []
    all_llm_tasks_rehab_followup: list[dict] = []
    all_llm_tasks_final_dx_proximity: list[dict] = []
    all_llm_tasks_gate1_rescore: list[dict] = []
    all_llm_tasks_gate2_rescore: list[dict] = []

    index_rows: list[dict[str, str]] = []
    parse_anomaly_rows: list[dict] = []
    review_index_rows: list[dict[str, str]] = []
    field_anomaly_rows: list[dict] = []

    # Column-level anomalies across all files (missing common columns / extra rare columns).
    field_anomaly_rows.extend(scan_column_anomalies(dataset_cols_map, common_cols))

    # Persisted human feedback extracted from previous review workbooks (optional).
    review_feedback_df = pd.DataFrame()
    try:
        fb_path = work_dir / "feedback" / "检查匹配_复核反馈.csv"
        if fb_path.exists():
            review_feedback_df = pd.read_csv(fb_path, encoding="utf-8-sig")
    except Exception:
        review_feedback_df = pd.DataFrame()

    # Human-in-the-loop overrides (persisted under work/latest/overrides)
    overrides_path = work_dir / "overrides" / "检查匹配_逐轮修正.csv"
    overrides_path.parent.mkdir(parents=True, exist_ok=True)
    auto_audit_path = work_dir / "overrides" / "检查匹配_逐轮修正_自动映射记录.csv"
    apply_review_feedback_to_override_template(overrides_path, review_feedback_df, audit_path=auto_audit_path)
    check_round_overrides = load_check_round_overrides_csv(overrides_path)
    applied_override_audits: list[pd.DataFrame] = []
    check_rounds_review_all: list[pd.DataFrame] = []

    for fileset in discovery.datasets:
        parse_before = len(parse_anomaly_rows)
        loaded_raw = load_dataset(fileset)

        # Data quality scan + in-memory fixes (never write back to data/).
        doc_fixed: dict[str, pd.DataFrame] = {}
        judge_fixed: dict[str, pd.DataFrame] = {}

        for sheet, df in loaded_raw.doc_sheets.items():
            for a in scan_sheet_anomalies("doc", sheet, df):
                a.update({"center": fileset.center, "model": fileset.model})
                parse_anomaly_rows.append(a)
            df2, changes = canonicalize_sheet_columns(df)
            for c in changes:
                details = {"new_column": c.get("new_column")}
                if c.get("action") == "merge_duplicate_after_rename":
                    details = {
                        "new_column": c.get("new_column"),
                        "columns": c.get("columns"),
                        "conflict_count": c.get("conflict_count"),
                        "conflict_cases": c.get("conflict_cases"),
                    }
                parse_anomaly_rows.append(
                    {
                        "center": fileset.center,
                        "model": fileset.model,
                        "kind": "doc",
                        "sheet": sheet,
                        "anomaly_type": f"canonicalize_{c.get('action')}",
                        "column": c.get("column"),
                        "details": json.dumps(details, ensure_ascii=False),
                    }
                )
            doc_fixed[sheet] = df2

        for sheet, df in loaded_raw.judge_sheets.items():
            for a in scan_sheet_anomalies("judge", sheet, df):
                a.update({"center": fileset.center, "model": fileset.model})
                parse_anomaly_rows.append(a)
            df2, changes = canonicalize_sheet_columns(df)
            for c in changes:
                details = {"new_column": c.get("new_column")}
                if c.get("action") == "merge_duplicate_after_rename":
                    details = {
                        "new_column": c.get("new_column"),
                        "columns": c.get("columns"),
                        "conflict_count": c.get("conflict_count"),
                        "conflict_cases": c.get("conflict_cases"),
                    }
                parse_anomaly_rows.append(
                    {
                        "center": fileset.center,
                        "model": fileset.model,
                        "kind": "judge",
                        "sheet": sheet,
                        "anomaly_type": f"canonicalize_{c.get('action')}",
                        "column": c.get("column"),
                        "details": json.dumps(details, ensure_ascii=False),
                    }
                )
            judge_fixed[sheet] = df2

        # Cross-sheet repairs (in-memory only; never write back to data/).
        doc_fixed, repaired_nested = repair_doc_decision_nested_fields(doc_fixed)
        for a in repaired_nested:
            a.update({"center": fileset.center, "model": fileset.model, "kind": "doc"})
            parse_anomaly_rows.append(a)

        doc_fixed, repaired = repair_doc_loop_decision_duplicate_raw_json(doc_fixed)
        for a in repaired:
            a.update({"center": fileset.center, "model": fileset.model, "kind": "doc"})
            parse_anomaly_rows.append(a)

        loaded = LoadedData(fileset=loaded_raw.fileset, gt=loaded_raw.gt, doc_sheets=doc_fixed, judge_sheets=judge_fixed)
        check_rounds = build_check_rounds(loaded, match_score_threshold=match_score_threshold)
        if not check_round_overrides.empty:
            check_rounds, applied_audit = apply_check_round_overrides(check_rounds, check_round_overrides)
            if applied_audit is not None and not applied_audit.empty:
                applied_override_audits.append(applied_audit)

        # Field-level missing value anomalies (doc/judge), based on stage outputs.
        field_anomaly_rows.extend(
            scan_missing_value_anomalies(fileset=fileset, doc_sheets=doc_fixed, judge_sheets=judge_fixed)
        )

        # Judge missing in loop rounds (executed but no judge data).
        if not check_rounds.empty:
            tmp_missing = check_rounds.copy()
            tmp_missing["ai_total_requested_count"] = pd.to_numeric(tmp_missing.get("ai_total_requested_count"), errors="coerce").fillna(0)
            judge_status_series = tmp_missing.get("judge_status", "").astype(str)
            missing_mask = (tmp_missing["ai_total_requested_count"] > 0) & pd.to_numeric(
                tmp_missing.get("matched_count_final"), errors="coerce"
            ).isna() & (~judge_status_series.str.contains("未评测", na=False))
            if missing_mask.any():
                for _, r in tmp_missing.loc[missing_mask].iterrows():
                    field_anomaly_rows.append(
                        {
                            "center": fileset.center,
                            "model": fileset.model,
                            "kind": "judge",
                            "sheet": r.get("stage"),
                            "case_id": r.get("case_id"),
                            "round_idx": r.get("round_idx"),
                            "anomaly_type": "judge_missing_check_round",
                            "column": "检查匹配(轮次)",
                            "details": json.dumps({"note": "AI有请求但裁判匹配信息缺失"}, ensure_ascii=False),
                        }
                    )

        # Detect doc-request missing cases (still computable via judge lists, but should be visible for human review).
        try:
            cr_tmp = check_rounds.copy()
            cr_tmp["ai_total_requested_count"] = pd.to_numeric(cr_tmp.get("ai_total_requested_count"), errors="coerce").fillna(0)
            cr_tmp["doc_request_col"] = cr_tmp.get("doc_request_col", "").astype(str)
            cr_tmp["case_id"] = cr_tmp.get("case_id", "").astype(str)

            # 1) Doc request text missing, judge lists used as proxy ("判官列表(兜底)")
            miss = cr_tmp.loc[
                (cr_tmp["stage"].isin(["D1_Outpatient_Loop", "D2_Admission_Loop"]))
                & (cr_tmp["ai_total_requested_count"] > 0)
                & (cr_tmp["doc_request_col"] == "判官列表(兜底)")
            ]
            if not miss.empty:
                for stage, g in miss.groupby("stage"):
                    cids = sorted(set(g["case_id"].tolist()), key=case_id_sort_key)
                    parse_anomaly_rows.append(
                        {
                            "center": fileset.center,
                            "model": fileset.model,
                            "kind": "doc",
                            "sheet": stage,
                            "anomaly_type": "doc_request_missing_used_judge_list",
                            "column": "doc_request_col",
                            "details": json.dumps({"count": len(cids), "case_ids": cids[:20]}, ensure_ascii=False),
                        }
                    )

            # 2) Extracted request column missing, recovered from round raw JSON ("第i轮_医生_原始JSON:需要补充检查")
            recovered = cr_tmp.loc[
                (cr_tmp["stage"].isin(["D1_Outpatient_Loop", "D2_Admission_Loop"]))
                & (cr_tmp["ai_total_requested_count"] > 0)
                & (cr_tmp["doc_request_col"].str.contains("原始JSON:", regex=False))
            ]
            if not recovered.empty:
                for stage, g in recovered.groupby("stage"):
                    cids = sorted(set(g["case_id"].tolist()), key=case_id_sort_key)
                    parse_anomaly_rows.append(
                        {
                            "center": fileset.center,
                            "model": fileset.model,
                            "kind": "doc",
                            "sheet": stage,
                            "anomaly_type": "doc_request_recovered_from_round_raw_json",
                            "column": "doc_request_col",
                            "details": json.dumps({"count": len(cids), "case_ids": cids[:20]}, ensure_ascii=False),
                        }
                    )
        except Exception:
            pass

        tables: MetricTables = build_metric_tables(loaded, check_rounds=check_rounds)

        # Field packs (for human review / Excel analysis)
        d1_decision_fields = build_d1_decision_fields(loaded)
        d2_decision_fields = build_d2_decision_fields(loaded)
        d3_decision_fields = build_d3_decision_fields(loaded)
        d4_rehab_fields = build_d4_rehab_fields(loaded)
        judge_scores = build_judge_scores_by_case(loaded)

        metrics_by_case = tables.metrics_by_case
        if not metrics_by_case.empty and not judge_scores.empty:
            metrics_by_case = metrics_by_case.merge(judge_scores, on=["center", "model", "case_id"], how="left")
        if not metrics_by_case.empty:
            metrics_by_case_full_rows.append(metrics_by_case.copy())
        if not tables.case_index.empty:
            case_index_rows.append(tables.case_index.copy())
        if not check_rounds.empty:
            check_rounds_rows.append(check_rounds.copy())
        if not tables.gt_checks.empty:
            gt_checks_rows.append(tables.gt_checks.copy())
        if not judge_scores.empty:
            judge_scores_rows.append(judge_scores.copy())
        if not d1_decision_fields.empty:
            d1_decision_fields_rows.append(d1_decision_fields.copy())
        if not d2_decision_fields.empty:
            d2_decision_fields_rows.append(d2_decision_fields.copy())
        if not d3_decision_fields.empty:
            d3_decision_fields_rows.append(d3_decision_fields.copy())
        if not d4_rehab_fields.empty:
            d4_rehab_fields_rows.append(d4_rehab_fields.copy())
        if not metrics_by_case.empty:
            keep_cols = [
                "center",
                "model",
                "case_id",
                "is_evaluated_inferred__D1_Outpatient_Loop",
                "D1_Outpatient_Loop__requested_unique_count",
                "D1_Outpatient_Loop__matched_unique_count",
                "D1_Outpatient_Loop__unmatched_unique_count",
                "D2_Check__requested_unique_count",
                "D2_Check__matched_unique_count",
                "D2_Check__unmatched_unique_count",
            ]
            metrics_by_case_rows.append(metrics_by_case.loc[:, [c for c in keep_cols if c in metrics_by_case.columns]].copy())

        check_rounds_audit_summary, check_rounds_review = _build_check_rounds_audit_tables(check_rounds)
        if check_rounds_audit_summary is not None and not check_rounds_audit_summary.empty:
            check_rounds_audit_rows.append(check_rounds_audit_summary.copy())
        if check_rounds_review is not None and not check_rounds_review.empty:
            check_rounds_review_all.append(check_rounds_review.copy())

        d1_decision_anomalies = pd.DataFrame()
        if not tables.case_index.empty and "is_anomaly_d1_to_d2" in tables.case_index.columns:
            d1_decision_anomalies = tables.case_index.loc[
                tables.case_index["is_anomaly_d1_to_d2"] == True, ["center", "model", "case_id"]
            ].copy()
        if not d1_decision_anomalies.empty:
            d1_decision_anomalies_rows.append(d1_decision_anomalies.copy())

        group_dir = out_dir / fileset.center / fileset.model
        group_dir.mkdir(parents=True, exist_ok=True)

        # Build enriched center-model summary (add judge score aggregates)
        judge_summary = pd.DataFrame()
        if not judge_scores.empty:
            js = judge_scores.copy()
            # Normalize boolean-like columns to 0/1 for aggregation.
            def to01(v):
                if v is None:
                    return pd.NA
                if isinstance(v, bool):
                    return 1 if v else 0
                s = str(v).strip().lower()
                if s in {"1", "true", "yes", "y"}:
                    return 1
                if s in {"0", "false", "no", "n"}:
                    return 0
                return pd.NA

            for c in ["gate1_continue", "gate2_continue", "d3_continue"]:
                if c in js.columns:
                    js[c] = js[c].map(to01)
            # Gate2 secondary review ("二审") if present.
            if "gate2_secondary_is_reasonable" in js.columns:
                js["gate2_secondary_is_reasonable"] = js["gate2_secondary_is_reasonable"].map(to01)
            for col in js.columns:
                if col in {"center", "model", "case_id"}:
                    continue
                js[col] = pd.to_numeric(js[col], errors="coerce")
            judge_summary = js.groupby(["center", "model"], as_index=False).agg(
                gate1_continue_rate=("gate1_continue", "mean"),
                gate1_continue_n=("gate1_continue", "count"),
                gate1_overall_score_mean=("gate1_overall_score", "mean"),
                gate1_dx_score_mean=("gate1_dx_score", "mean"),
                gate1_check_score_mean=("gate1_check_score", "mean"),
                gate1_check_match_degree_mean=("gate1_check_match_degree", "mean"),
                gate2_continue_rate=("gate2_continue", "mean"),
                gate2_continue_n=("gate2_continue", "count"),
                gate2_secondary_rate=("gate2_secondary_is_reasonable", "mean"),
                gate2_secondary_n=("gate2_secondary_is_reasonable", "count"),
                gate2_overall_score_mean=("gate2_overall_score", "mean"),
                gate2_revised_dx_score_mean=("gate2_revised_dx_score", "mean"),
                gate2_surgery_score_mean=("gate2_surgery_score", "mean"),
                d3_continue_rate=("d3_continue", "mean"),
                d3_continue_n=("d3_continue", "count"),
                d3_overall_score_mean=("d3_overall_score", "mean"),
                d3_dx_score_mean=("d3_dx_score", "mean"),
                d3_plan_score_mean=("d3_plan_score", "mean"),
                d4_overall_score_mean=("d4_overall_score", "mean"),
                d4_rehab_score_mean=("d4_rehab_score", "mean"),
                d4_followup_score_mean=("d4_followup_score", "mean"),
            )

        metrics_by_center_model = tables.metrics_by_center_model.copy()
        if not metrics_by_center_model.empty and not judge_summary.empty:
            metrics_by_center_model = metrics_by_center_model.merge(judge_summary, on=["center", "model"], how="left")

        # Human review workbook (Chinese)
        parse_anomalies_df = pd.DataFrame(parse_anomaly_rows) if parse_anomaly_rows else pd.DataFrame()
        review_spec = build_review_workbook_spec(
            center=fileset.center,
            model=fileset.model,
            metrics_by_center_model=metrics_by_center_model,
            metrics_by_case=metrics_by_case,
            judge_scores_by_case=judge_scores,
            check_rounds=check_rounds,
            check_rounds_review=check_rounds_review,
            parse_anomalies=parse_anomalies_df,
            field_anomalies=pd.DataFrame(field_anomaly_rows) if field_anomaly_rows else pd.DataFrame(),
            review_feedback=review_feedback_df,
        )
        review_xlsx_path = group_dir / "审阅表.xlsx"
        try:
            export_review_workbook_xlsx(review_xlsx_path, review_spec)
        except PermissionError:
            # If the file is open/locked, write to a new filename.
            review_xlsx_path = group_dir / "审阅表_新.xlsx"
            export_review_workbook_xlsx(review_xlsx_path, review_spec)

        # Optional: keep the full raw workbook for debugging (writes many columns/sheets)
        if export_excel:
            raw_dir = work_dir / "_raw" / fileset.center / fileset.model
            raw_dir.mkdir(parents=True, exist_ok=True)
            export_metrics_source_data_xlsx(
                raw_dir / "metrics_source_data.xlsx",
                sheets=[
                    ExcelSheetSpec("case_index", tables.case_index),
                    ExcelSheetSpec("check_rounds", check_rounds),
                    ExcelSheetSpec("check_rounds_audit_summary", check_rounds_audit_summary),
                    ExcelSheetSpec("check_rounds_review", check_rounds_review),
                    ExcelSheetSpec("d1_decision_anomalies", d1_decision_anomalies),
                    ExcelSheetSpec("gt_checks", tables.gt_checks),
                    ExcelSheetSpec("judge_scores_by_case", judge_scores),
                    ExcelSheetSpec("d1_decision_fields", d1_decision_fields),
                    ExcelSheetSpec("d2_decision_fields", d2_decision_fields),
                    ExcelSheetSpec("d3_decision_fields", d3_decision_fields),
                    ExcelSheetSpec("d4_rehab_fields", d4_rehab_fields),
                    ExcelSheetSpec("metrics_by_case", metrics_by_case),
                    ExcelSheetSpec("metrics_by_center_model", metrics_by_center_model),
                ],
            )

        # LLM task generation (no API call here; just prepare tasks JSONL).
        allowed_case_ids: set[str] | None = None
        if llm_task_scope == "review_only":
            allowed_case_ids = set()
            if not check_rounds_review.empty and "case_id" in check_rounds_review.columns:
                allowed_case_ids |= set(check_rounds_review["case_id"].astype(str))
            if not d1_decision_anomalies.empty and "case_id" in d1_decision_anomalies.columns:
                allowed_case_ids |= set(d1_decision_anomalies["case_id"].astype(str))
            if not metrics_by_case.empty and "issues" in metrics_by_case.columns:
                allowed_case_ids |= set(metrics_by_case.loc[metrics_by_case["issues"].astype(str).str.strip().ne(""), "case_id"].astype(str))
            if allowed_case_ids == set():
                allowed_case_ids = None
        elif llm_task_scope == "all":
            allowed_case_ids = None
        else:
            raise ValueError(f"Unknown llm_task_scope={llm_task_scope!r}. Expected 'review_only' or 'all'.")

        if "llm.diagnosis_semantic_match" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_dx_topk_d1.extend(
                generate_d1_diagnosis_topk_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.rationale_quality" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_rationale.extend(
                generate_rationale_quality_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.fact_consistency_and_missing" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_fact_consistency.extend(
                generate_fact_consistency_and_missing_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.memory_retention" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_memory_retention.extend(
                generate_memory_retention_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.cross_stage_consistency" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_cross_stage.extend(
                generate_cross_stage_consistency_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.final_diagnosis_proximity" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_final_dx_proximity.extend(
                generate_final_diagnosis_proximity_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.gate1_dx_rescore" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_gate1_rescore.extend(
                generate_gate1_dx_rescore_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.gate2_dx_plan_rescore" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_gate2_rescore.extend(
                generate_gate2_dx_plan_rescore_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.unmatched_check_reasonableness" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_unmatched_checks.extend(
                generate_unmatched_check_reasonableness_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.diagnosis_quality" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_dx_quality.extend(
                generate_diagnosis_quality_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.plan_quality" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_plan_quality.extend(
                generate_plan_quality_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.diagnosis_bias" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_dx_bias.extend(
                generate_diagnosis_bias_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )
        if "llm.rehab_followup_quality" in LLM_TASK_ALLOWLIST:
            all_llm_tasks_rehab_followup.extend(
                generate_rehab_followup_quality_tasks(project_root, loaded, allowed_case_ids=allowed_case_ids)
            )

        metrics_by_center_model_rows.append(metrics_by_center_model)
        index_rows.append(
            {
                "center": fileset.center,
                "model": fileset.model,
                "group_dir": str(group_dir.relative_to(out_dir)).replace("\\", "/"),
            }
        )
        review_index_rows.append(
            {
                "中心": fileset.center,
                "模型": fileset.model,
                "审阅表": str(review_xlsx_path.relative_to(out_dir)).replace("\\", "/"),
                "需要复核_检查匹配(行数)": int(len(check_rounds_review)) if check_rounds_review is not None else 0,
                "状态疑点(行数)": int(len(review_spec.status_issues)) if review_spec.status_issues is not None else 0,
                "解析异常(条数)": int(len(parse_anomaly_rows) - parse_before),
            }
        )

    # Global summary (small)
    metrics_by_center_model = pd.concat(metrics_by_center_model_rows, ignore_index=True) if metrics_by_center_model_rows else pd.DataFrame()
    metrics_by_case_all = pd.concat(metrics_by_case_rows, ignore_index=True) if metrics_by_case_rows else pd.DataFrame()
    metrics_by_case_full_all = (
        pd.concat(metrics_by_case_full_rows, ignore_index=True) if metrics_by_case_full_rows else pd.DataFrame()
    )
    case_index_all = pd.concat(case_index_rows, ignore_index=True) if case_index_rows else pd.DataFrame()
    check_rounds_all = pd.concat(check_rounds_rows, ignore_index=True) if check_rounds_rows else pd.DataFrame()
    check_rounds_audit_all = (
        pd.concat(check_rounds_audit_rows, ignore_index=True) if check_rounds_audit_rows else pd.DataFrame()
    )
    check_rounds_review_all_df = (
        pd.concat(check_rounds_review_all, ignore_index=True) if check_rounds_review_all else pd.DataFrame()
    )
    d1_decision_anomalies_all = (
        pd.concat(d1_decision_anomalies_rows, ignore_index=True) if d1_decision_anomalies_rows else pd.DataFrame()
    )
    gt_checks_all = pd.concat(gt_checks_rows, ignore_index=True) if gt_checks_rows else pd.DataFrame()
    judge_scores_all = pd.concat(judge_scores_rows, ignore_index=True) if judge_scores_rows else pd.DataFrame()
    d1_decision_fields_all = (
        pd.concat(d1_decision_fields_rows, ignore_index=True) if d1_decision_fields_rows else pd.DataFrame()
    )
    d2_decision_fields_all = (
        pd.concat(d2_decision_fields_rows, ignore_index=True) if d2_decision_fields_rows else pd.DataFrame()
    )
    d3_decision_fields_all = (
        pd.concat(d3_decision_fields_rows, ignore_index=True) if d3_decision_fields_rows else pd.DataFrame()
    )
    d4_rehab_fields_all = (
        pd.concat(d4_rehab_fields_rows, ignore_index=True) if d4_rehab_fields_rows else pd.DataFrame()
    )
    if out_dir.name != "latest":
        metrics_by_center_model.to_csv(summary_dir / "metrics_by_center_model_raw.csv", index=False, encoding="utf-8-sig")

    if export_excel:
        export_metrics_source_data_xlsx(
            out_dir / "metrics_source_data.xlsx",
            sheets=[
                ExcelSheetSpec("case_index", case_index_all),
                ExcelSheetSpec("check_rounds", check_rounds_all),
                ExcelSheetSpec("check_rounds_audit_summary", check_rounds_audit_all),
                ExcelSheetSpec("check_rounds_review", check_rounds_review_all_df),
                ExcelSheetSpec("d1_decision_anomalies", d1_decision_anomalies_all),
                ExcelSheetSpec("gt_checks", gt_checks_all),
                ExcelSheetSpec("judge_scores_by_case", judge_scores_all),
                ExcelSheetSpec("d1_decision_fields", d1_decision_fields_all),
                ExcelSheetSpec("d2_decision_fields", d2_decision_fields_all),
                ExcelSheetSpec("d3_decision_fields", d3_decision_fields_all),
                ExcelSheetSpec("d4_rehab_fields", d4_rehab_fields_all),
                ExcelSheetSpec("metrics_by_case", metrics_by_case_full_all),
                ExcelSheetSpec("metrics_by_center_model", metrics_by_center_model),
            ],
        )

    # Update (upsert) override template for human editing.
    if check_rounds_review_all:
        review_all = pd.concat(check_rounds_review_all, ignore_index=True)
        update_check_round_override_template(overrides_path, review_all)

    # Export override audit (what was actually applied in this run)
    applied_overrides_df = pd.concat(applied_override_audits, ignore_index=True) if applied_override_audits else pd.DataFrame()
    if not applied_overrides_df.empty:
        (work_dir / "overrides").mkdir(parents=True, exist_ok=True)
        applied_overrides_df.to_csv(work_dir / "overrides" / "检查匹配_逐轮修正_应用记录.csv", index=False, encoding="utf-8-sig")

    # Write LLM task files for this run BEFORE trying to match cached results.
    # (run_id=latest clears work/latest/llm_tasks at start, so we must recreate it here)
    if all_llm_tasks_dx_topk_d1:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.diagnosis_semantic_match.D1.jsonl", all_llm_tasks_dx_topk_d1)
    if all_llm_tasks_rationale:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.rationale_quality.jsonl", all_llm_tasks_rationale)
    if all_llm_tasks_fact_consistency:
        write_tasks_jsonl(
            work_dir / "llm_tasks" / "llm.fact_consistency_and_missing.jsonl",
            all_llm_tasks_fact_consistency,
        )
    if all_llm_tasks_memory_retention:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.memory_retention.jsonl", all_llm_tasks_memory_retention)
    if all_llm_tasks_cross_stage:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.cross_stage_consistency.jsonl", all_llm_tasks_cross_stage)
    if all_llm_tasks_final_dx_proximity:
        write_tasks_jsonl(
            work_dir / "llm_tasks" / "llm.final_diagnosis_proximity.jsonl",
            all_llm_tasks_final_dx_proximity,
        )
    if all_llm_tasks_gate1_rescore:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.gate1_dx_rescore.jsonl", all_llm_tasks_gate1_rescore)
    if all_llm_tasks_gate2_rescore:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.gate2_dx_plan_rescore.jsonl", all_llm_tasks_gate2_rescore)
    if all_llm_tasks_unmatched_checks:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.unmatched_check_reasonableness.jsonl", all_llm_tasks_unmatched_checks)
    if all_llm_tasks_dx_quality:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.diagnosis_quality.jsonl", all_llm_tasks_dx_quality)
    if all_llm_tasks_plan_quality:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.plan_quality.jsonl", all_llm_tasks_plan_quality)
    if all_llm_tasks_dx_bias:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.diagnosis_bias.jsonl", all_llm_tasks_dx_bias)
    if all_llm_tasks_rehab_followup:
        write_tasks_jsonl(work_dir / "llm_tasks" / "llm.rehab_followup_quality.jsonl", all_llm_tasks_rehab_followup)

    # --- Optional: integrate cached LLM results into summary tables (no API calls)
    def _inputs_hash_task(t: dict) -> str:
        payload = {
            "task": t.get("task"),
            "task_id": t.get("task_id"),
            "prompt_template": t.get("prompt_template"),
            "inputs": t.get("inputs"),
        }
        return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()

    def _load_task_keys(task_jsonl: Path) -> set[tuple[str, str]]:
        keys: set[tuple[str, str]] = set()
        if not task_jsonl.exists():
            return keys
        with task_jsonl.open("r", encoding="utf-8") as f:
            for ln in f:
                ln = ln.strip()
                if not ln:
                    continue
                try:
                    t = json.loads(ln)
                except Exception:
                    continue
                task_id = str(t.get("task_id", "")).strip()
                if not task_id:
                    continue
                keys.add((task_id, _inputs_hash_task(t)))
        return keys

    def _pick_llm_results_jsonl() -> Path | None:
        llm_dir = work_dir / "llm_results"
        if not llm_dir.exists():
            return None
        preferred = llm_dir / "gemini-3-pro-preview-thinking__gala_api.jsonl"
        if preferred.exists():
            return preferred
        cands = list(llm_dir.glob("*.jsonl"))
        if not cands:
            return None
        return max(cands, key=lambda p: p.stat().st_mtime)

    results_jsonl = _pick_llm_results_jsonl()
    tasks_dir = work_dir / "llm_tasks"
    llm_check_reason_by_case: pd.DataFrame | None = None
    llm_topk_by_case: pd.DataFrame | None = None
    llm_check_reason_detail: pd.DataFrame | None = None
    llm_topk_detail: pd.DataFrame | None = None
    if results_jsonl and results_jsonl.exists() and tasks_dir.exists() and not metrics_by_case_all.empty:
        unmatched_keys = _load_task_keys(tasks_dir / "llm.unmatched_check_reasonableness.jsonl")
        dx_keys = _load_task_keys(tasks_dir / "llm.diagnosis_semantic_match.D1.jsonl")

        unmatched_rows: list[dict] = []
        dx_rows: list[dict] = []
        unmatched_detail_rows: list[dict] = []
        dx_detail_rows: list[dict] = []

        # Keep only the latest record per (task_id, inputs_hash) to avoid duplication across reruns.
        latest: dict[tuple[str, str], dict] = {}
        with results_jsonl.open("r", encoding="utf-8", errors="ignore") as f:
            for ln in f:
                ln = ln.strip()
                if not ln:
                    continue
                try:
                    r = json.loads(ln)
                except Exception:
                    continue
                task_id = str(r.get("task_id", "")).strip()
                ih = str(r.get("inputs_hash", "")).strip()
                if not task_id or not ih:
                    continue
                latest[(task_id, ih)] = r

        for (task_id, ih), r in latest.items():
            task = str(r.get("task", "")).strip()
            if r.get("error"):
                continue
            parsed = r.get("response_json") if isinstance(r.get("response_json"), dict) else None
            if not parsed:
                continue

            if task == "llm.unmatched_check_reasonableness" and (task_id, ih) in unmatched_keys:
                items = parsed.get("items") or []
                all_items: list[str] = []
                meaningful_items: list[str] = []
                redundant_items: list[str] = []
                meaningful_non_redundant: list[str] = []
                per_item_notes: list[dict] = []
                for it in items:
                    if not isinstance(it, dict):
                        continue
                    chk = str(it.get("check_item") or "").strip()
                    if not chk:
                        continue
                    all_items.append(chk)
                    meaningful = int(str(it.get("is_clinically_meaningful", 0) or 0)) == 1
                    redundant = int(str(it.get("is_redundant", 0) or 0)) == 1
                    if meaningful:
                        meaningful_items.append(chk)
                    if redundant:
                        redundant_items.append(chk)
                    if meaningful and not redundant:
                        meaningful_non_redundant.append(chk)
                    per_item_notes.append(
                        {
                            "check_item": chk,
                            "is_clinically_meaningful": it.get("is_clinically_meaningful"),
                            "is_redundant": it.get("is_redundant"),
                            "expected_impact": it.get("expected_impact"),
                            "rationale": it.get("rationale"),
                            "confidence": it.get("confidence"),
                        }
                    )

                unmatched_rows.append(
                    {
                        "center": r.get("center"),
                        "model": r.get("doc_model"),
                        "case_id": r.get("case_id"),
                        "stage": r.get("stage"),
                        "llm_unmatched_items_total": len(items),
                        "llm_unmatched_meaningful_count": len(meaningful_non_redundant),
                        "llm_unmatched_redundant_count": len(redundant_items),
                    }
                )

                inputs = r.get("inputs") if isinstance(r.get("inputs"), dict) else {}
                unmatched_detail_rows.append(
                    {
                        "center": r.get("center"),
                        "model": r.get("doc_model"),
                        "case_id": r.get("case_id"),
                        "stage": r.get("stage"),
                        "task_id": task_id,
                        "prompt_template": r.get("prompt_template"),
                        "prompt_sha256": (inputs or {}).get("prompt_sha256"),
                        "prompt": r.get("prompt"),
                        "input__unmatched_checks_to_judge": (inputs or {}).get("unmatched_checks_to_judge"),
                        "input__ai_requested_checks_all": (inputs or {}).get("ai_requested_checks_all"),
                        "input__judge_matched_checks_all": (inputs or {}).get("judge_matched_checks_all"),
                        "input__judge_unmatched_checks_all": (inputs or {}).get("judge_unmatched_checks_all"),
                        "input__prior_matched_checks": (inputs or {}).get("prior_matched_checks"),
                        "input__gt_checks_text": (inputs or {}).get("gt_checks_text"),
                        "input__gt_facts": (inputs or {}).get("gt_facts"),
                        "input__ai_rationale": (inputs or {}).get("ai_rationale"),
                        "output__meaningful_unmatched_checks": "\n".join(meaningful_non_redundant),
                        "output__redundant_or_not_meaningful_checks": "\n".join(
                            sorted(set(redundant_items) | set([x for x in all_items if x not in set(meaningful_items)]))
                        ),
                        "output__items_json": json.dumps(per_item_notes, ensure_ascii=False),
                        "response_json": json.dumps(parsed or {}, ensure_ascii=False),
                    }
                )

            elif task == "llm.diagnosis_semantic_match" and (task_id, ih) in dx_keys:
                by_k: dict[int, dict] = {}
                for item in parsed.get("results") or []:
                    if not isinstance(item, dict):
                        continue
                    try:
                        k = int(item.get("k"))
                    except Exception:
                        continue
                    by_k[k] = item
                dx_rows.append(
                    {
                        "center": r.get("center"),
                        "model": r.get("doc_model"),
                        "case_id": r.get("case_id"),
                        "stage": r.get("stage"),
                        "topk_hit1": (by_k.get(1) or {}).get("hit_any", (by_k.get(1) or {}).get("hit")),
                        "topk_hit3": (by_k.get(3) or {}).get("hit_any", (by_k.get(3) or {}).get("hit")),
                        "topk_hit5": (by_k.get(5) or {}).get("hit_any", (by_k.get(5) or {}).get("hit")),
                        "topk_weighted1": (by_k.get(1) or {}).get("weighted_score"),
                        "topk_weighted3": (by_k.get(3) or {}).get("weighted_score"),
                        "topk_weighted5": (by_k.get(5) or {}).get("weighted_score"),
                    }
                )

                inputs = r.get("inputs") if isinstance(r.get("inputs"), dict) else {}
                dx_detail_rows.append(
                    {
                        "center": r.get("center"),
                        "model": r.get("doc_model"),
                        "case_id": r.get("case_id"),
                        "stage": r.get("stage"),
                        "task_id": task_id,
                        "prompt_template": r.get("prompt_template"),
                        "prompt_sha256": (inputs or {}).get("prompt_sha256"),
                        "prompt": r.get("prompt"),
                        "input__gt_diagnosis_text": (inputs or {}).get("gt_diagnosis_text"),
                        "input__ai_diagnosis_list_text": (inputs or {}).get("ai_diagnosis_list_text"),
                        "input__strictness": json.dumps((inputs or {}).get("strictness") or {}, ensure_ascii=False),
                        "output__results_json": json.dumps(parsed.get("results") or [], ensure_ascii=False),
                        "output__notes_json": json.dumps(parsed.get("notes") or [], ensure_ascii=False),
                        "response_json": json.dumps(parsed or {}, ensure_ascii=False),
                    }
                )

        # ---- unmatched check reasonableness -> extended (match + reasonable) rate (unique-count view)
        if unmatched_rows:
            llm_unmatched = pd.DataFrame(unmatched_rows)
            d1 = llm_unmatched.loc[llm_unmatched["stage"] == "D1_Outpatient"].rename(
                columns={
                    "llm_unmatched_items_total": "D1_LLM未匹配项数",
                    "llm_unmatched_meaningful_count": "D1_LLM合理项数",
                    "llm_unmatched_redundant_count": "D1_LLM冗余项数",
                }
            )
            d1 = d1.drop(columns=["stage"], errors="ignore")
            d2 = llm_unmatched.loc[llm_unmatched["stage"] == "D2_Admission"].rename(
                columns={
                    "llm_unmatched_items_total": "D2_LLM未匹配项数",
                    "llm_unmatched_meaningful_count": "D2_LLM合理项数",
                    "llm_unmatched_redundant_count": "D2_LLM冗余项数",
                }
            )
            d2 = d2.drop(columns=["stage"], errors="ignore")

            merged = metrics_by_case_all.merge(d1, on=["center", "model", "case_id"], how="left").merge(
                d2, on=["center", "model", "case_id"], how="left"
            )

            def _safe_rate(matched, meaningful, requested):
                if requested is None or requested == 0:
                    return pd.NA
                num = (matched or 0) + (meaningful or 0)
                num = min(num, requested)
                return num / requested

            merged["D1检查合理性率(匹配+合理，唯一)"] = pd.NA
            has_d1_llm = merged["D1_LLM未匹配项数"].notna()
            den = pd.to_numeric(merged.get("D1_Outpatient_Loop__requested_unique_count"), errors="coerce")
            num_matched = pd.to_numeric(merged.get("D1_Outpatient_Loop__matched_unique_count"), errors="coerce")
            num_meaningful = pd.to_numeric(merged.get("D1_LLM合理项数"), errors="coerce")
            # Baseline: if there is no unmatched item, this equals strict match rate (meaningful=0).
            merged.loc[(den > 0) & num_matched.notna(), "D1检查合理性率(匹配+合理，唯一)"] = (
                num_matched.fillna(0).clip(upper=den) / den
            )
            # If LLM judged unmatched items, add meaningful items into numerator (cap at denominator).
            merged.loc[has_d1_llm & (den > 0), "D1检查合理性率(匹配+合理，唯一)"] = (
                (num_matched.fillna(0) + num_meaningful.fillna(0)).clip(upper=den) / den
            )

            merged["D2检查合理性率(匹配+合理，唯一)"] = pd.NA
            has_d2_llm = merged["D2_LLM未匹配项数"].notna()
            den2 = pd.to_numeric(merged.get("D2_Check__requested_unique_count"), errors="coerce")
            num_matched2 = pd.to_numeric(merged.get("D2_Check__matched_unique_count"), errors="coerce")
            num_meaningful2 = pd.to_numeric(merged.get("D2_LLM合理项数"), errors="coerce")
            merged.loc[(den2 > 0) & num_matched2.notna(), "D2检查合理性率(匹配+合理，唯一)"] = (
                num_matched2.fillna(0).clip(upper=den2) / den2
            )
            merged.loc[has_d2_llm & (den2 > 0), "D2检查合理性率(匹配+合理，唯一)"] = (
                (num_matched2.fillna(0) + num_meaningful2.fillna(0)).clip(upper=den2) / den2
            )

            # Export per-case (small) table for review
            out_case = merged[
                [
                    "center",
                    "model",
                    "case_id",
                    "D1_Outpatient_Loop__requested_unique_count",
                    "D1_Outpatient_Loop__matched_unique_count",
                    "D1_LLM未匹配项数",
                    "D1_LLM合理项数",
                    "D1检查合理性率(匹配+合理，唯一)",
                    "D2_Check__requested_unique_count",
                    "D2_Check__matched_unique_count",
                    "D2_LLM未匹配项数",
                    "D2_LLM合理项数",
                    "D2检查合理性率(匹配+合理，唯一)",
                ]
            ].copy()
            out_case = out_case.rename(
                columns={
                    "center": "中心",
                    "model": "模型",
                    "case_id": "病例ID",
                    "D1_Outpatient_Loop__requested_unique_count": "D1_请求检查数(唯一)",
                    "D1_Outpatient_Loop__matched_unique_count": "D1_匹配数(唯一)",
                    "D2_Check__requested_unique_count": "D2_请求检查数(唯一)",
                    "D2_Check__matched_unique_count": "D2_匹配数(唯一)",
                }
            )
            out_case["_case_sort"] = out_case["病例ID"].map(case_id_sort_key)
            out_case = out_case.sort_values(by=["中心", "模型", "_case_sort"], kind="mergesort").drop(columns=["_case_sort"])
            out_case.to_csv(summary_dir / "检查合理性(LLM)_按病例.csv", index=False, encoding="utf-8-sig")
            llm_check_reason_by_case = out_case
            if unmatched_detail_rows:
                det = pd.DataFrame(unmatched_detail_rows)
                det["_case_sort"] = det["case_id"].map(case_id_sort_key)
                det = det.sort_values(by=["center", "model", "_case_sort", "stage"], kind="mergesort").drop(columns=["_case_sort"])
                llm_check_reason_detail = det.rename(
                    columns={
                        "center": "中心",
                        "model": "模型",
                        "case_id": "病例ID",
                        "stage": "阶段",
                        "task_id": "任务ID",
                        "prompt_template": "提示词模板",
                        "prompt_sha256": "提示词版本(hash)",
                        "prompt": "提示词(完整)",
                        "input__unmatched_checks_to_judge": "输入_需判定合理性的未匹配检查",
                        "input__ai_requested_checks_all": "输入_AI提出检查(全量)",
                        "input__judge_matched_checks_all": "输入_裁判判定匹配检查(全量)",
                        "input__judge_unmatched_checks_all": "输入_裁判判定未匹配检查(全量)",
                        "input__gt_checks_text": "输入_GT已执行检查原文",
                        "input__gt_facts": "输入_GT病历事实",
                        "input__ai_rationale": "输入_AI理由/思维",
                        "output__meaningful_unmatched_checks": "输出_合理的未匹配检查(列表)",
                        "output__redundant_or_not_meaningful_checks": "输出_不合理/冗余未匹配检查(列表)",
                        "output__items_json": "输出_逐项判定(JSON)",
                        "response_json": "输出_JSON原文",
                    }
                )

            # Aggregate per center/model and merge into metrics_by_center_model
            llm_cm = out_case.groupby(["中心", "模型"], as_index=False).agg(
                d1_check_match_plus_reasonable_unique_mean=("D1检查合理性率(匹配+合理，唯一)", "mean"),
                d1_check_match_plus_reasonable_unique_n=("D1检查合理性率(匹配+合理，唯一)", "count"),
                d1_check_reason_llm_covered_n=("D1_LLM未匹配项数", "count"),
                d2_check_match_plus_reasonable_unique_mean=("D2检查合理性率(匹配+合理，唯一)", "mean"),
                d2_check_match_plus_reasonable_unique_n=("D2检查合理性率(匹配+合理，唯一)", "count"),
                d2_check_reason_llm_covered_n=("D2_LLM未匹配项数", "count"),
            )
            llm_cm = llm_cm.rename(columns={"中心": "center", "模型": "model"})
            metrics_by_center_model = metrics_by_center_model.merge(llm_cm, on=["center", "model"], how="left")

        # ---- diagnosis TopK -> add per-center/model aggregates (optional, small)
        if dx_rows:
            dx_df = pd.DataFrame(dx_rows)
            for c in ["topk_hit1", "topk_hit3", "topk_hit5", "topk_weighted1", "topk_weighted3", "topk_weighted5"]:
                dx_df[c] = pd.to_numeric(dx_df.get(c), errors="coerce")
            dx_cm = dx_df.groupby(["center", "model"], as_index=False).agg(
                d1_topk_hit1_rate=("topk_hit1", "mean"),
                d1_topk_hit3_rate=("topk_hit3", "mean"),
                d1_topk_hit5_rate=("topk_hit5", "mean"),
                d1_topk_weighted1_mean=("topk_weighted1", "mean"),
                d1_topk_weighted3_mean=("topk_weighted3", "mean"),
                d1_topk_weighted5_mean=("topk_weighted5", "mean"),
                d1_topk_n=("case_id", "count"),
            )
            dx_case_cn = dx_df.rename(
                columns={
                    "center": "中心",
                    "model": "模型",
                    "case_id": "病例ID",
                    "stage": "阶段",
                    "topk_hit1": "k=1命中(任一GT)(0/1)",
                    "topk_hit3": "k=3命中(任一GT)(0/1)",
                    "topk_hit5": "k=5命中(任一GT)(0/1)",
                    "topk_weighted1": "k=1加权覆盖得分(0-1)",
                    "topk_weighted3": "k=3加权覆盖得分(0-1)",
                    "topk_weighted5": "k=5加权覆盖得分(0-1)",
                }
            )
            dx_case_cn["_case_sort"] = dx_case_cn["病例ID"].map(case_id_sort_key)
            dx_case_cn = dx_case_cn.sort_values(by=["中心", "模型", "_case_sort", "阶段"], kind="mergesort").drop(columns=["_case_sort"])
            dx_case_cn.to_csv(summary_dir / "诊断TopK(LLM)_按病例.csv", index=False, encoding="utf-8-sig")
            llm_topk_by_case = dx_case_cn
            if dx_detail_rows:
                det = pd.DataFrame(dx_detail_rows)
                det["_case_sort"] = det["case_id"].map(case_id_sort_key)
                det = det.sort_values(by=["center", "model", "_case_sort", "stage"], kind="mergesort").drop(columns=["_case_sort"])
                llm_topk_detail = det.rename(
                    columns={
                        "center": "中心",
                        "model": "模型",
                        "case_id": "病例ID",
                        "stage": "阶段",
                        "task_id": "任务ID",
                        "prompt_template": "提示词模板",
                        "prompt_sha256": "提示词版本(hash)",
                        "prompt": "提示词(完整)",
                        "input__gt_diagnosis_text": "输入_GT诊断文本",
                        "input__ai_diagnosis_list_text": "输入_AI诊断列表",
                        "input__strictness": "输入_严格度(JSON)",
                        "output__results_json": "输出_results(JSON)",
                        "output__notes_json": "输出_notes(JSON)",
                        "response_json": "输出_JSON原文",
                    }
                )
            metrics_by_center_model = metrics_by_center_model.merge(dx_cm, on=["center", "model"], how="left")

    # Chinese-friendly summary outputs (what you should actually open)
    cn_map = {
        "center": "中心",
        "model": "模型",
        "cases": "病例数",
        "d1_check_match_rate_mean": "D1检查匹配率(病例均值)",
        "d1_check_match_rate_weighted": "D1检查匹配率(总匹配/总请求)",
        "d1_check_match_rate_n": "D1检查匹配率样本数(已评测)",
        "d1_check_match_rate_requested_n": "D1检查匹配率样本数(有请求)",
        "d1_check_match_rate_not_scored_n": "D1检查匹配率未评测数(裁判缺失)",
        "d2_check_match_rate_mean": "D2检查匹配率(病例均值)",
        "d2_check_match_rate_weighted": "D2检查匹配率(总匹配/总请求)",
        "d2_check_match_rate_n": "D2检查匹配率样本数(已评测)",
        "d2_check_match_rate_requested_n": "D2检查匹配率样本数(有请求)",
        "d2_check_match_rate_not_scored_n": "D2检查匹配率未评测数(裁判缺失)",
        "d1_inefficiency_by_zero_mean": "D1无效循环率(0匹配)均值",
        "d1_inefficiency_by_zero_n": "D1无效循环率样本数",
        "d2_inefficiency_by_zero_mean": "D2无效循环率(0匹配)均值",
        "d2_inefficiency_by_zero_n": "D2无效循环率样本数",
        "d1_check_match_plus_reasonable_unique_mean": "D1检查合理性率均值(匹配+合理/请求，唯一)",
        "d1_check_match_plus_reasonable_unique_n": "D1检查合理性率样本数(有请求)",
        "d1_check_reason_llm_covered_n": "D1检查合理性率LLM覆盖数",
        "d2_check_match_plus_reasonable_unique_mean": "D2检查合理性率均值(匹配+合理/请求，唯一)",
        "d2_check_match_plus_reasonable_unique_n": "D2检查合理性率样本数(有请求)",
        "d2_check_reason_llm_covered_n": "D2检查合理性率LLM覆盖数",
        "d1_topk_hit1_rate": "D1 Top1语义命中率(任一GT)",
        "d1_topk_hit3_rate": "D1 Top3语义命中率(任一GT)",
        "d1_topk_hit5_rate": "D1 Top5语义命中率(任一GT)",
        "d1_topk_weighted1_mean": "D1 Top1加权覆盖得分均值",
        "d1_topk_weighted3_mean": "D1 Top3加权覆盖得分均值",
        "d1_topk_weighted5_mean": "D1 Top5加权覆盖得分均值",
        "d1_topk_n": "D1 TopK样本数(LLM)",
        # Gate (flow-based, from inferred status)
        "gate1_flow_pass_rate": "Gate1 流程通过率(基于流程状态)",
        "gate1_flow_reached_n": "Gate1 流程样本数(到达D1决策)",
        "gate1_flow_pass_n": "Gate1 流程通过数(到达D2及以后)",
        "gate2_flow_pass_rate": "Gate2 流程通过率(基于流程状态)",
        "gate2_flow_reached_n": "Gate2 流程样本数(到达D2决策)",
        "gate2_flow_pass_n": "Gate2 流程通过数(到达D3及以后)",
        # Gate (judge-based)
        "gate1_continue_rate": "Gate1 Judge通过率(继续评测)",
        "gate1_continue_n": "Gate1 Judge样本数",
        "gate1_overall_score_mean": "Gate1 Judge总分均值",
        "gate1_dx_score_mean": "Gate1 Judge诊断评分均值",
        "gate1_check_score_mean": "Gate1 Judge检查匹配评分均值",
        "gate1_check_match_degree_mean": "Gate1 Judge检查匹配度均值",
        "gate2_continue_rate": "Gate2 Judge初审通过率(继续评测)",
        "gate2_continue_n": "Gate2 Judge初审样本数",
        "gate2_secondary_rate": "Gate2 Judge复审通过率(is_reasonable)",
        "gate2_secondary_n": "Gate2 Judge复审样本数",
        "gate2_overall_score_mean": "Gate2 Judge总分均值",
        "gate2_revised_dx_score_mean": "Gate2 Judge修正诊断评分均值",
        "gate2_surgery_score_mean": "Gate2 Judge手术方案评分均值",
        "d3_continue_rate": "D3是否继续评测通过率",
        "d3_continue_n": "D3是否继续评测样本数",
        "d3_overall_score_mean": "D3总分均值",
        "d3_dx_score_mean": "D3最终诊断评分均值",
        "d3_plan_score_mean": "D3术后方案评分均值",
        "d4_overall_score_mean": "D4总分均值",
        "d4_rehab_score_mean": "D4康复计划评分均值",
        "d4_followup_score_mean": "D4随访计划评分均值",
    }
    metrics_cn = metrics_by_center_model.rename(columns=cn_map).copy() if not metrics_by_center_model.empty else pd.DataFrame()
    if not metrics_cn.empty:
        metrics_cn.to_csv(summary_dir / "指标汇总_中心模型.csv", index=False, encoding="utf-8-sig")

    if parse_anomaly_rows:
        pa_csv = pd.DataFrame(parse_anomaly_rows).copy()
        pa_csv = pa_csv.rename(
            columns={
                "center": "中心",
                "model": "模型",
                "kind": "来源(医生/裁判)",
                "sheet": "工作表",
                "anomaly_type": "异常类型",
                "column": "字段/列名",
                "details": "详情",
            }
        )
        if "来源(医生/裁判)" in pa_csv.columns:
            pa_csv["来源(医生/裁判)"] = pa_csv["来源(医生/裁判)"].map({"doc": "医生", "judge": "裁判"}).fillna(pa_csv["来源(医生/裁判)"])
        if "异常类型" in pa_csv.columns:
            pa_csv.insert(pa_csv.columns.get_loc("异常类型") + 1, "异常类型代码", pa_csv["异常类型"])
            pa_csv["异常类型"] = (
                pa_csv["异常类型"]
                .map(
                    {
                        "header_known_mojibake": "表头乱码(已自动修复)",
                        "header_non_ascii_latin": "表头包含异常拉丁字符(已自动修复)",
                        "check_col_contains_decision_terms": "检查列疑似混入决策内容",
                        "loop_raw_json_equals_decision_raw_json": "循环原始JSON与对应决策完全相同(已自动按未经过处理)",
                        "doc_request_missing_used_judge_list": "AI请求文本缺失(已用裁判列表兜底)",
                        "doc_request_recovered_from_round_raw_json": "AI请求列缺失(已从轮次原始JSON恢复)",
                        "canonicalize_rename": "表头已自动修正(重命名)",
                        "canonicalize_merge_duplicate_after_rename": "表头重复(重命名后)已自动合并",
                        "canonicalize_drop_duplicate_after_rename": "表头重复(重命名后)已自动丢弃",
                    }
                )
                .fillna(pa_csv["异常类型"])
            )
        pa_csv.to_csv(summary_dir / "解析异常.csv", index=False, encoding="utf-8-sig")

    # Doc/Judge 字段异常清单 + 交互修复入口（不改 data/）
    field_anomaly_df = pd.DataFrame(field_anomaly_rows) if field_anomaly_rows else pd.DataFrame()
    if not field_anomaly_df.empty:
        # Natural sort by case_id for readability
        field_anomaly_df["_case_sort"] = field_anomaly_df.get("case_id", "").map(case_id_sort_key)
        field_anomaly_df["round_idx"] = pd.to_numeric(field_anomaly_df.get("round_idx"), errors="coerce")
        field_anomaly_df = field_anomaly_df.sort_values(
            by=["center", "model", "_case_sort", "sheet", "round_idx"], kind="mergesort"
        ).drop(columns=["_case_sort"])
        field_anomaly_df.to_csv(summary_dir / "doc_judge_字段异常清单.csv", index=False, encoding="utf-8-sig")

        repairs_dir = work_dir / "repairs"
        repairs_dir.mkdir(parents=True, exist_ok=True)

        missing_mask = field_anomaly_df["anomaly_type"].isin(
            [
                "doc_missing_required_field",
                "doc_missing_required_any",
                "judge_missing_required_field",
                "judge_missing_all_fields",
                "judge_missing_check_round",
            ]
        )
        missing_df = field_anomaly_df.loc[missing_mask].copy()
        if not missing_df.empty:
            missing_df = missing_df.rename(
                columns={
                    "center": "中心",
                    "model": "模型",
                    "kind": "来源(医生/裁判)",
                    "sheet": "工作表",
                    "case_id": "病例ID",
                    "round_idx": "轮次",
                    "anomaly_type": "异常类型",
                    "column": "字段/列名",
                    "details": "详情",
                }
            )
            missing_df["来源(医生/裁判)"] = missing_df["来源(医生/裁判)"].map({"doc": "医生", "judge": "裁判"}).fillna(
                missing_df["来源(医生/裁判)"]
            )
            missing_df["修正值"] = ""
            missing_df["采用修正(0/1)"] = ""
            missing_df["备注"] = ""
            missing_df = missing_df[
                [
                    "中心",
                    "模型",
                    "来源(医生/裁判)",
                    "工作表",
                    "病例ID",
                    "轮次",
                    "字段/列名",
                    "异常类型",
                    "详情",
                    "修正值",
                    "采用修正(0/1)",
                    "备注",
                ]
            ]
            missing_df.to_csv(repairs_dir / "缺失字段_修复模板.csv", index=False, encoding="utf-8-sig")

        col_mask = field_anomaly_df["anomaly_type"].isin(["missing_common_column", "extra_rare_column"])
        col_df = field_anomaly_df.loc[col_mask].copy()
        if not col_df.empty:
            col_df = col_df.rename(
                columns={
                    "center": "中心",
                    "model": "模型",
                    "kind": "来源(医生/裁判)",
                    "sheet": "工作表",
                    "anomaly_type": "异常类型",
                    "column": "字段/列名",
                    "details": "详情",
                }
            )
            col_df["来源(医生/裁判)"] = col_df["来源(医生/裁判)"].map({"doc": "医生", "judge": "裁判"}).fillna(
                col_df["来源(医生/裁判)"]
            )
            col_df["建议动作(drop/rename/merge)"] = ""
            col_df["修正列名"] = ""
            col_df["采用修正(0/1)"] = ""
            col_df["备注"] = ""
            col_df = col_df[
                [
                    "中心",
                    "模型",
                    "来源(医生/裁判)",
                    "工作表",
                    "字段/列名",
                    "异常类型",
                    "详情",
                    "建议动作(drop/rename/merge)",
                    "修正列名",
                    "采用修正(0/1)",
                    "备注",
                ]
            ]
            col_df.to_csv(repairs_dir / "列异常_修复模板.csv", index=False, encoding="utf-8-sig")

    if index_rows and out_dir.name != "latest":
        pd.DataFrame(index_rows).to_csv(summary_dir / "index.csv", index=False, encoding="utf-8-sig")
    if review_index_rows:
        pd.DataFrame(review_index_rows).to_csv(summary_dir / "审阅表索引.csv", index=False, encoding="utf-8-sig")

        # One-row-per-center-model checklist (easy to tick)
        checklist_df = pd.DataFrame(review_index_rows).copy()
        if not checklist_df.empty:
            checklist_df.insert(checklist_df.shape[1], "已复核(请勾选)", "")
            checklist_df.insert(checklist_df.shape[1], "处理人", "")
            checklist_df.insert(checklist_df.shape[1], "备注", "")
            checklist_df.to_csv(summary_dir / "问题清单.csv", index=False, encoding="utf-8-sig")
        # Per user request: do NOT generate summary Markdown docs here.

    # Single summary workbook (easy to open)
    summary_xlsx = summary_dir / "总览.xlsx"
    sheets = []
    if not metrics_cn.empty:
        sheets.append(ExcelSheetSpec("指标汇总_中心模型", metrics_cn))
    if llm_check_reason_by_case is not None and not llm_check_reason_by_case.empty:
        sheets.append(ExcelSheetSpec("检查合理性(LLM)_按病例", llm_check_reason_by_case))
    if llm_check_reason_detail is not None and not llm_check_reason_detail.empty:
        sheets.append(ExcelSheetSpec("检查合理性(LLM)_明细", llm_check_reason_detail))
    if llm_topk_by_case is not None and not llm_topk_by_case.empty:
        sheets.append(ExcelSheetSpec("诊断TopK(LLM)_按病例", llm_topk_by_case))
    if llm_topk_detail is not None and not llm_topk_detail.empty:
        sheets.append(ExcelSheetSpec("诊断TopK(LLM)_明细", llm_topk_detail))
    if parse_anomaly_rows:
        pa = pd.DataFrame(parse_anomaly_rows).copy().rename(
            columns={
                "center": "中心",
                "model": "模型",
                "kind": "来源(医生/裁判)",
                "sheet": "工作表",
                "anomaly_type": "异常类型",
                "column": "字段/列名",
                "details": "详情",
            }
        )
        if "来源(医生/裁判)" in pa.columns:
            pa["来源(医生/裁判)"] = pa["来源(医生/裁判)"].map({"doc": "医生", "judge": "裁判"}).fillna(pa["来源(医生/裁判)"])
        if "异常类型" in pa.columns:
            pa.insert(pa.columns.get_loc("异常类型") + 1, "异常类型代码", pa["异常类型"])
            pa["异常类型"] = (
                pa["异常类型"]
                .map(
                    {
                        "header_known_mojibake": "表头乱码(已知模式)",
                        "header_non_ascii_latin": "表头包含异常拉丁字符",
                        "check_col_contains_decision_terms": "检查列疑似混入决策内容",
                        "canonicalize_rename": "表头已自动修正(重命名)",
                        "canonicalize_drop_duplicate_after_rename": "表头重复(重命名后)已自动丢弃",
                    }
                )
                .fillna(pa["异常类型"])
            )
        sheets.append(ExcelSheetSpec("解析异常", pa))
    if not field_anomaly_df.empty:
        fa_doc = field_anomaly_df.loc[field_anomaly_df["kind"] == "doc"].copy()
        fa_judge = field_anomaly_df.loc[field_anomaly_df["kind"] == "judge"].copy()
        for fa in [fa_doc, fa_judge]:
            if fa.empty:
                continue
            fa.rename(
                columns={
                    "center": "中心",
                    "model": "模型",
                    "kind": "来源(医生/裁判)",
                    "sheet": "工作表",
                    "case_id": "病例ID",
                    "round_idx": "轮次",
                    "anomaly_type": "异常类型",
                    "column": "字段/列名",
                    "details": "详情",
                },
                inplace=True,
            )
            fa["来源(医生/裁判)"] = fa["来源(医生/裁判)"].map({"doc": "医生", "judge": "裁判"}).fillna(
                fa["来源(医生/裁判)"]
            )
            if "异常类型" in fa.columns:
                fa.insert(fa.columns.get_loc("异常类型") + 1, "异常类型代码", fa["异常类型"])
                fa["异常类型"] = (
                    fa["异常类型"]
                    .map(
                        {
                            "missing_common_column": "缺失通用列",
                            "extra_rare_column": "非通用列(可能异常)",
                            "doc_missing_required_field": "医生关键字段缺失",
                            "doc_missing_required_any": "医生关键字段缺失(需至少一项)",
                            "judge_missing_required_field": "裁判关键字段缺失",
                            "judge_missing_all_fields": "裁判关键字段全缺失",
                            "judge_missing_check_round": "裁判缺失/未评测(检查匹配)",
                        }
                    )
                    .fillna(fa["异常类型"])
                )
        if not fa_doc.empty:
            sheets.append(ExcelSheetSpec("doc_异常清单", fa_doc))
        if not fa_judge.empty:
            sheets.append(ExcelSheetSpec("judge_异常清单", fa_judge))
    if review_index_rows:
        sheets.append(ExcelSheetSpec("审阅表索引", pd.DataFrame(review_index_rows)))
    if sheets:
        export_metrics_source_data_xlsx(summary_xlsx, sheets=sheets)

    # Append LLM-derived tables into each center×model review workbook (for easier auditing).
    if review_index_rows and (
        llm_check_reason_by_case is not None
        or llm_check_reason_detail is not None
        or llm_topk_by_case is not None
        or llm_topk_detail is not None
    ):
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except Exception:
            load_workbook = None  # type: ignore[assignment]

        def _append_sheet(xlsx_path: Path, sheet_name: str, df: pd.DataFrame) -> None:
            if df is None or df.empty:
                return
            if not xlsx_path.exists():
                return
            try:
                with pd.ExcelWriter(xlsx_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:  # type: ignore[arg-type]
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
            except TypeError:
                # Older pandas: remove sheet first, then append.
                if load_workbook is None:
                    return
                wb = load_workbook(xlsx_path)
                if sheet_name in wb.sheetnames:
                    wb.remove(wb[sheet_name])
                    wb.save(xlsx_path)
                with pd.ExcelWriter(xlsx_path, engine="openpyxl", mode="a") as writer:
                    df.to_excel(writer, sheet_name=sheet_name, index=False)

            if load_workbook is None:
                return
            wb = load_workbook(xlsx_path)
            ws = wb[sheet_name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            header_fill = PatternFill("solid", fgColor="1F4E79")
            header_font = Font(color="FFFFFF", bold=True)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center")

            # Column widths (sampled)
            max_row = min(ws.max_row, 600)
            for col_idx in range(1, ws.max_column + 1):
                letter = get_column_letter(col_idx)
                header_len = len(str(ws.cell(row=1, column=col_idx).value or ""))
                sample_len = 0
                for row_idx in range(2, max_row + 1):
                    v = ws.cell(row=row_idx, column=col_idx).value
                    if v is None:
                        continue
                    sample_len = max(sample_len, len(str(v)))
                ws.column_dimensions[letter].width = max(10, min(60, max(header_len, sample_len) + 2))

            wb.save(xlsx_path)

        for r in review_index_rows:
            center = str(r.get("中心", "")).strip()
            model = str(r.get("模型", "")).strip()
            rel = str(r.get("审阅表", "")).strip().replace("/", "\\")
            if not center or not model or not rel:
                continue
            xlsx_path = out_dir / rel
            if llm_check_reason_by_case is not None and {"中心", "模型"} <= set(llm_check_reason_by_case.columns):
                df = llm_check_reason_by_case.loc[
                    (llm_check_reason_by_case["中心"] == center) & (llm_check_reason_by_case["模型"] == model)
                ].copy()
                _append_sheet(xlsx_path, "8_LLM-检查合理性(按病例)", df)
            if llm_check_reason_detail is not None and {"中心", "模型"} <= set(llm_check_reason_detail.columns):
                df = llm_check_reason_detail.loc[
                    (llm_check_reason_detail["中心"] == center) & (llm_check_reason_detail["模型"] == model)
                ].copy()
                _append_sheet(xlsx_path, "8_LLM-检查合理性(明细)", df)
            if llm_topk_by_case is not None and {"中心", "模型"} <= set(llm_topk_by_case.columns):
                df = llm_topk_by_case.loc[(llm_topk_by_case["中心"] == center) & (llm_topk_by_case["模型"] == model)].copy()
                _append_sheet(xlsx_path, "8_LLM-诊断TopK(D1)", df)
            if llm_topk_detail is not None and {"中心", "模型"} <= set(llm_topk_detail.columns):
                df = llm_topk_detail.loc[(llm_topk_detail["中心"] == center) & (llm_topk_detail["模型"] == model)].copy()
                _append_sheet(xlsx_path, "8_LLM-诊断TopK(明细)", df)


def _update_latest_outputs_index(outputs_root: Path, out_dir: Path) -> None:
    """
    Convenience pointer to avoid hunting timestamped folders.
    Writes:
    - outputs/LATEST.txt
    - outputs/LATEST.md
    """
    run_id = out_dir.name
    (outputs_root / "LATEST.txt").write_text(run_id + "\n", encoding="utf-8")

    idx_path = out_dir / "summary" / "index.csv"
    lines: list[str] = []
    lines.append("# Latest Outputs")
    lines.append("")
    lines.append(f"- run_id: `{run_id}`")
    lines.append(f"- path: `{out_dir}`")
    lines.append("")

    if idx_path.exists():
        try:
            df = pd.read_csv(idx_path)
        except Exception:
            df = pd.DataFrame()
        if not df.empty and {"center", "model", "group_dir"} <= set(df.columns):
            lines.append("## Groups (center × model)")
            for _, r in df.iterrows():
                group_dir = str(r.get("group_dir", "")).strip().replace("\\", "/")
                if not group_dir:
                    continue
                review_xlsx = f"outputs/{run_id}/{group_dir}/审阅表.xlsx"
                lines.append(f"- {r.get('center')} | {r.get('model')}: `{review_xlsx}`")
            lines.append("")

    (outputs_root / "LATEST.md").write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
