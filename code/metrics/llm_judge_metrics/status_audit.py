from __future__ import annotations

import csv
import json
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .discovery import discover_datasets
from .ingest import SHEETS
from .parsing import is_empty
from .quality import canonicalize_sheet_columns
from .repairs import repair_doc_decision_nested_fields
from .sorting import case_id_sort_key
from .status import STAGE_ORDER, detect_d1_decision_anomaly, infer_flow_status_row, _row_has_output_for_stage


_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
_FILL_ORANGE = PatternFill("solid", fgColor="FFEB9C")
_FILL_GREY = PatternFill("solid", fgColor="E7E6E6")
_FILL_RED = PatternFill("solid", fgColor="FFC7CE")

SPECIAL_STATUS = "终止于门诊决策（特殊情况）"
SPECIAL_D2_STATUS = "D1特殊"


_KIND_CN = {"doc": "医生", "judge": "裁判"}
_SHEET_CN = {
    "D1_Outpatient_Loop": "D1门诊循环",
    "D1_Outpatient_Decision": "D1门诊决策",
    "D2_Admission_Loop": "D2入院循环",
    "D2_Admission_Decision": "D2入院决策",
    "D3_Surgery_Decision": "D3手术决策",
    "D4_Rehab_Plan": "D4康复计划",
}


def _make_run_id(tag: str) -> str:
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    safe_tag = "".join(ch if ch.isalnum() or ch in {"-", "_"} else "-" for ch in tag.strip().lower()).strip("-")
    safe_tag = safe_tag or "status-audit"
    if safe_tag in {"latest", "current"}:
        return "latest"
    return f"{ts}_{safe_tag}"


@dataclass(frozen=True)
class StatusDiffRow:
    center: str
    model: str
    kind: str  # doc | judge
    sheet: str
    case_id: str
    status_before: str
    status_after: str


def _read_status_map(path: Path, sheet: str) -> dict[str, str]:
    df = pd.read_excel(path, sheet_name=sheet, engine="openpyxl", usecols=[0, 1])
    if df.shape[1] < 2:
        return {}
    cid_col = df.columns[0]
    status_col = df.columns[1]
    m: dict[str, str] = {}
    for _, r in df.iterrows():
        cid = str(r.get(cid_col, "")).strip()
        if not cid:
            continue
        m[cid] = str(r.get(status_col, "")).strip()
    return m


def _style_status_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Header style
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    # Conditional formatting: apply to column B ("状态") if present.
    if ws.max_column < 2:
        return
    header_b = ws.cell(row=1, column=2).value
    if not header_b or "状态" not in str(header_b):
        return

    col = get_column_letter(2)
    rng = f"{col}2:{col}{ws.max_row}"
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=[f'ISNUMBER(SEARCH("特殊情况",{col}2))'], fill=_FILL_RED, stopIfTrue=True),
    )
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("顺利通过",{col}2))'], fill=_FILL_GREEN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("终止",{col}2))'], fill=_FILL_ORANGE))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("未经过",{col}2))'], fill=_FILL_GREY))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("Anomaly",{col}2))'], fill=_FILL_RED))


def _boolish(v: Any) -> bool | None:
    if is_empty(v):
        return None
    if isinstance(v, bool):
        return bool(v)
    s = str(v).strip().lower()
    if s in {"true", "1", "yes", "y", "是", "通过"}:
        return True
    if s in {"false", "0", "no", "n", "否", "不通过"}:
        return False
    return None


def _has_any_cols(df: pd.DataFrame, row: pd.Series | None, keys: list[str]) -> bool:
    if row is None or df is None or df.empty:
        return False
    cols = [c for c in df.columns if any(k in str(c) for k in keys)]
    return any(cols) and any(not is_empty(row.get(c)) for c in cols)


def _infer_status_maps_for_doc(path: Path) -> tuple[pd.DataFrame, dict[str, dict[str, str]]]:
    doc_sheets: dict[str, pd.DataFrame] = {}
    for sheet in STAGE_ORDER:
        doc_sheets[sheet] = pd.read_excel(path, sheet_name=sheet, engine="openpyxl")

    # Fix known misplacements for inference only (D1 decision accidentally stored in D2 decision).
    try:
        doc_d1 = doc_sheets.get("D1_Outpatient_Decision")
        doc_d2 = doc_sheets.get("D2_Admission_Decision")
        misplaced = _detect_misplaced_d1_to_d2_cases(doc_d1, doc_d2)
        if misplaced:
            doc_d1, doc_d2 = _restore_misplaced_d1_from_d2_df(doc_d1, doc_d2, misplaced)
            doc_sheets["D1_Outpatient_Decision"] = doc_d1
            doc_sheets["D2_Admission_Decision"] = doc_d2
    except Exception:
        pass

    # Prefer a union of case_ids across sheets instead of anchoring to one sheet,
    # because some files may have missing/shifted rows in a single sheet.
    case_ids_set: set[str] = set()
    for df in doc_sheets.values():
        if df is None or df.empty:
            continue
        cid_col = df.columns[0]
        for x in df[cid_col].dropna().astype(str).tolist():
            cid = str(x).strip()
            if cid:
                case_ids_set.add(cid)

    if not case_ids_set:
        return pd.DataFrame(), {}

    case_ids = sorted(case_ids_set)

    rows: list[dict] = []
    per_sheet: dict[str, dict[str, str]] = {s: {} for s in STAGE_ORDER}
    for case_id in case_ids:
        r = infer_flow_status_row(doc_sheets, case_id)
        rows.append(r)
        for sheet in STAGE_ORDER:
            per_sheet[sheet][case_id] = str(r.get(f"doc_status_inferred__{sheet}", "未经过"))
    return pd.DataFrame(rows), per_sheet


def _apply_status_map_to_workbook(wb, status_map: dict[str, dict[str, str]]) -> list[tuple[str, str, str]]:
    """
    Apply inferred status to each sheet's column B, return list of (sheet, case_id, old_status).
    """
    changed: list[tuple[str, str, str]] = []
    for sheet in SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        _style_status_sheet(ws)
        m = status_map.get(sheet, {})
        for row_idx in range(2, ws.max_row + 1):
            cid = ws.cell(row=row_idx, column=1).value
            if cid is None:
                continue
            case_id = str(cid).strip()
            if not case_id:
                continue
            new_status = m.get(case_id)
            if new_status is None:
                continue
            old_status = str(ws.cell(row=row_idx, column=2).value or "").strip()
            if old_status != new_status:
                ws.cell(row=row_idx, column=2).value = new_status
                changed.append((sheet, case_id, old_status))
    return changed


def _apply_special_cases_to_doc_decision(ws, special_cases: set[str]) -> None:
    if ws is None or not special_cases:
        return
    for row_idx in range(2, ws.max_row + 1):
        cid = ws.cell(row=row_idx, column=1).value
        if cid is None:
            continue
        case_id = str(cid).strip()
        if not case_id or case_id not in special_cases:
            continue
        # Mark status and clear other fields
        ws.cell(row=row_idx, column=2).value = SPECIAL_STATUS
        ws.cell(row=row_idx, column=2).fill = _FILL_RED
        for col_idx in range(3, ws.max_column + 1):
            ws.cell(row=row_idx, column=col_idx).value = None


def _apply_special_cases_to_d2_decision(ws, special_cases: set[str]) -> None:
    if ws is None or not special_cases:
        return
    for row_idx in range(2, ws.max_row + 1):
        cid = ws.cell(row=row_idx, column=1).value
        if cid is None:
            continue
        case_id = str(cid).strip()
        if not case_id or case_id not in special_cases:
            continue
        # Mark status only (do not clear other fields)
        ws.cell(row=row_idx, column=2).value = SPECIAL_D2_STATUS
        ws.cell(row=row_idx, column=2).fill = _FILL_RED


def _detect_misplaced_d1_to_d2_cases(doc_d1: pd.DataFrame, doc_d2: pd.DataFrame) -> set[str]:
    """
    Detect cases that were mistakenly moved from D1 decision into D2 decision.
    Heuristic: D1 has no outputs, but D2 contains D1-like fields or D1-specific keys in raw JSON.
    """
    if doc_d2 is None or doc_d2.empty:
        return set()
    if doc_d1 is None or doc_d1.empty:
        doc_d1 = pd.DataFrame(columns=[doc_d2.columns[0]])

    case_col = doc_d2.columns[0]
    cols_need_more = [c for c in doc_d2.columns if "需要进一步检查" in str(c)]
    cols_need_out = [c for c in doc_d2.columns if "需要补充门诊检查" in str(c)]

    d1_like_keys = [
        "初步诊断列表",
        "建议检查项目",
        "初步诊断思维",
        "建议检查思维",
        "诊断置信度",
        "检查方案置信度",
        "需要进一步检查",
        "需要补充门诊检查",
        "门诊信息汇总",
    ]
    d2_like_keys = [
        "修正诊断",
        "修正诊断思维",
        "初步治疗方案",
        "治疗方案思维",
        "诊疗经过回顾",
        "手术方案",
    ]
    raw_cols = [c for c in doc_d2.columns if str(c).strip() == "医生决策_原始JSON"]

    def _raw_has_d1_keys(v: Any) -> bool:
        if is_empty(v):
            return False
        try:
            obj = json.loads(str(v))
        except Exception:
            s = str(v)
            return any(k in s for k in d1_like_keys)
        if isinstance(obj, dict):
            return any(k in obj for k in d1_like_keys)
        return False

    def _raw_has_d2_keys(v: Any) -> bool:
        if is_empty(v):
            return False
        try:
            obj = json.loads(str(v))
        except Exception:
            s = str(v)
            return any(k in s for k in d2_like_keys)
        if isinstance(obj, dict):
            return any(k in obj for k in d2_like_keys)
        return False

    misplaced: set[str] = set()
    for _, row in doc_d2.iterrows():
        case_id = str(row.get(case_col, "")).strip()
        if not case_id:
            continue
        d1_row = None
        if doc_d1 is not None and not doc_d1.empty:
            sel = doc_d1.loc[doc_d1[doc_d1.columns[0]].astype(str) == case_id]
            if not sel.empty:
                d1_row = sel.iloc[0]
        has_d1_output = _row_has_output_for_stage("D1_Outpatient_Decision", doc_d1, d1_row)

        if has_d1_output:
            continue

        has_d1_like = _has_any_cols(doc_d2, row, d1_like_keys)
        has_d2_like = _has_any_cols(doc_d2, row, d2_like_keys)
        if raw_cols:
            for c in raw_cols:
                if _raw_has_d1_keys(row.get(c)):
                    has_d1_like = True
                if _raw_has_d2_keys(row.get(c)):
                    has_d2_like = True
        if has_d2_like:
            # D2 decision content exists; do not treat as misplaced D1.
            continue

        need_more = None
        for c in cols_need_more:
            b = _boolish(row.get(c))
            if b is not None:
                need_more = b
                break
        need_out = None
        for c in cols_need_out:
            b = _boolish(row.get(c))
            if b is not None:
                need_out = b
                break

        if has_d1_like or need_more is True or need_out is True:
            misplaced.add(case_id)

    return misplaced


def _restore_misplaced_d1_from_d2_df(
    doc_d1: pd.DataFrame, doc_d2: pd.DataFrame, misplaced_cases: set[str]
) -> tuple[pd.DataFrame, pd.DataFrame]:
    if not misplaced_cases or doc_d2 is None or doc_d2.empty:
        return doc_d1, doc_d2
    if doc_d1 is None or doc_d1.empty:
        doc_d1 = pd.DataFrame(columns=list(doc_d2.columns))

    d1 = doc_d1.copy()
    d2 = doc_d2.copy()

    d1_cid = d1.columns[0]
    d2_cid = d2.columns[0]
    d1_status_col = d1.columns[1] if d1.shape[1] > 1 else None
    d2_status_col = d2.columns[1] if d2.shape[1] > 1 else None

    def _extract_from_raw(raw: Any) -> dict[str, Any]:
        if is_empty(raw):
            return {}
        try:
            obj = json.loads(str(raw))
        except Exception:
            return {}
        if not isinstance(obj, dict):
            return {}
        out: dict[str, Any] = {}
        mapping = {
            "需要补充门诊检查": "医生决策_原始JSON_需要补充门诊检查",
            "需要进一步检查": "医生决策_原始JSON_需要进一步检查",
            "门诊信息汇总": "医生决策_原始JSON_门诊信息汇总",
            "初步诊断列表": "医生决策_原始JSON_初步诊断列表",
            "建议检查项目": "医生决策_原始JSON_建议检查项目",
            "初步诊断思维": "医生决策_原始JSON_初步诊断思维",
            "建议检查思维": "医生决策_原始JSON_建议检查思维",
        }
        for k, col in mapping.items():
            if k in obj and not is_empty(obj.get(k)):
                out[col] = obj.get(k)
        conf = obj.get("置信度评估")
        if isinstance(conf, dict):
            if not is_empty(conf.get("诊断置信度")):
                out["医生决策_原始JSON_置信度评估_诊断置信度"] = conf.get("诊断置信度")
            if not is_empty(conf.get("检查方案置信度")):
                out["医生决策_原始JSON_置信度评估_检查方案置信度"] = conf.get("检查方案置信度")
        return out

    def _normalize(v: Any) -> Any:
        if isinstance(v, (dict, list)):
            return json.dumps(v, ensure_ascii=False)
        return v

    for case_id in misplaced_cases:
        d2_row = d2.loc[d2[d2_cid].astype(str) == str(case_id)]
        if d2_row.empty:
            continue
        src = d2_row.iloc[0]
        raw_json = src.get("医生决策_原始JSON") if "医生决策_原始JSON" in d2.columns else None
        extracted = _extract_from_raw(raw_json)

        d1_row = d1.loc[d1[d1_cid].astype(str) == str(case_id)]
        if d1_row.empty:
            d1 = pd.concat([d1, pd.DataFrame([{d1_cid: case_id}])], ignore_index=True)
            d1_row = d1.loc[d1[d1_cid].astype(str) == str(case_id)]
        idx = d1_row.index[0]

        # Fill D1 raw JSON and parsed D1 fields (normal decision columns only).
        for col in d1.columns:
            if col == d1_cid:
                continue
            if col == "医生决策_原始JSON" and not is_empty(raw_json):
                if is_empty(d1.at[idx, col]):
                    d1.at[idx, col] = raw_json
                continue
            if col in extracted and is_empty(d1.at[idx, col]):
                d1.at[idx, col] = _normalize(extracted.get(col))

        # Clear D1 special columns that belong to D2 decision semantics.
        for col in d1.columns:
            col_s = str(col)
            if "修正诊断" in col_s or "初步治疗方案" in col_s or "治疗方案思维" in col_s or "诊疗经过回顾" in col_s:
                d1.at[idx, col] = pd.NA
            if "置信度评估_治疗方案置信度" in col_s:
                d1.at[idx, col] = pd.NA

        if d1_status_col is not None:
            d1.at[idx, d1_status_col] = "顺利通过"

        # Clear D2 row (keep case_id and status)
        d2_idx = d2_row.index[0]
        if d2_status_col is not None:
            d2.at[d2_idx, d2_status_col] = "未经过"
        for col in d2.columns[2:]:
            d2.at[d2_idx, col] = pd.NA

    return d1, d2


def _restore_misplaced_d1_from_d2(wb, misplaced_cases: set[str]) -> list[dict[str, Any]]:
    if not misplaced_cases:
        return []
    if "D1_Outpatient_Decision" not in wb.sheetnames or "D2_Admission_Decision" not in wb.sheetnames:
        return []
    ws_d1 = wb["D1_Outpatient_Decision"]
    ws_d2 = wb["D2_Admission_Decision"]

    # Build header maps
    d1_headers = [ws_d1.cell(row=1, column=c).value for c in range(1, ws_d1.max_column + 1)]
    d2_headers = [ws_d2.cell(row=1, column=c).value for c in range(1, ws_d2.max_column + 1)]
    d1_map = {str(h): i + 1 for i, h in enumerate(d1_headers) if h is not None}
    d2_map = {str(h): i + 1 for i, h in enumerate(d2_headers) if h is not None}

    def _extract_from_raw(raw: Any) -> dict[str, Any]:
        if is_empty(raw):
            return {}
        try:
            obj = json.loads(str(raw))
        except Exception:
            return {}
        if not isinstance(obj, dict):
            return {}
        out: dict[str, Any] = {}
        mapping = {
            "需要补充门诊检查": "医生决策_原始JSON_需要补充门诊检查",
            "需要进一步检查": "医生决策_原始JSON_需要进一步检查",
            "门诊信息汇总": "医生决策_原始JSON_门诊信息汇总",
            "初步诊断列表": "医生决策_原始JSON_初步诊断列表",
            "建议检查项目": "医生决策_原始JSON_建议检查项目",
            "初步诊断思维": "医生决策_原始JSON_初步诊断思维",
            "建议检查思维": "医生决策_原始JSON_建议检查思维",
        }
        for k, col in mapping.items():
            if k in obj and not is_empty(obj.get(k)):
                out[col] = obj.get(k)
        conf = obj.get("置信度评估")
        if isinstance(conf, dict):
            if not is_empty(conf.get("诊断置信度")):
                out["医生决策_原始JSON_置信度评估_诊断置信度"] = conf.get("诊断置信度")
            if not is_empty(conf.get("检查方案置信度")):
                out["医生决策_原始JSON_置信度评估_检查方案置信度"] = conf.get("检查方案置信度")
        return out

    def _normalize(v: Any) -> Any:
        if isinstance(v, (dict, list)):
            return json.dumps(v, ensure_ascii=False)
        return v

    audit: list[dict[str, Any]] = []
    for row_idx in range(2, ws_d2.max_row + 1):
        cid = ws_d2.cell(row=row_idx, column=1).value
        if cid is None:
            continue
        case_id = str(cid).strip()
        if case_id not in misplaced_cases:
            continue

        # Ensure D1 row exists
        d1_row_idx = None
        for r in range(2, ws_d1.max_row + 1):
            if str(ws_d1.cell(row=r, column=1).value or "").strip() == case_id:
                d1_row_idx = r
                break
        if d1_row_idx is None:
            d1_row_idx = ws_d1.max_row + 1
            ws_d1.cell(row=d1_row_idx, column=1).value = case_id

        raw_val = None
        raw_col = d2_map.get("医生决策_原始JSON")
        if raw_col is not None:
            raw_val = ws_d2.cell(row=row_idx, column=raw_col).value
        extracted = _extract_from_raw(raw_val)

        # Fill D1 raw JSON and parsed D1 fields only.
        for name, d1_col in d1_map.items():
            if name in {"病例ID", "case_id"}:
                continue
            if name == "医生决策_原始JSON":
                if raw_val is not None and is_empty(ws_d1.cell(row=d1_row_idx, column=d1_col).value):
                    ws_d1.cell(row=d1_row_idx, column=d1_col).value = raw_val
                continue
            if name in extracted and is_empty(ws_d1.cell(row=d1_row_idx, column=d1_col).value):
                ws_d1.cell(row=d1_row_idx, column=d1_col).value = _normalize(extracted.get(name))

        # Clear D1 special columns (D2 semantics)
        for name, d1_col in d1_map.items():
            name_s = str(name)
            if "修正诊断" in name_s or "初步治疗方案" in name_s or "治疗方案思维" in name_s or "诊疗经过回顾" in name_s:
                ws_d1.cell(row=d1_row_idx, column=d1_col).value = None
            if "置信度评估_治疗方案置信度" in name_s:
                ws_d1.cell(row=d1_row_idx, column=d1_col).value = None

        # Set D1 status to normal
        if 2 in range(1, ws_d1.max_column + 1):
            ws_d1.cell(row=d1_row_idx, column=2).value = "顺利通过"

        # Clear D2 row (keep case_id and status)
        if 2 in range(1, ws_d2.max_column + 1):
            ws_d2.cell(row=row_idx, column=2).value = "未经过"
        for col in range(3, ws_d2.max_column + 1):
            ws_d2.cell(row=row_idx, column=col).value = None

        audit.append({"case_id": case_id, "action": "restore_d1_from_d2"})

    return audit


def _rename_treatment_key_in_json(value: Any) -> Any:
    if is_empty(value):
        return value
    if isinstance(value, dict):
        if "治疗方案" in value and "初步治疗方案" not in value:
            value = dict(value)
            value["初步治疗方案"] = value.pop("治疗方案")
        return value
    s = str(value).strip()
    if not s:
        return value
    try:
        obj = json.loads(s)
    except Exception:
        return s.replace("\"治疗方案\"", "\"初步治疗方案\"")
    if isinstance(obj, dict):
        if "治疗方案" in obj and "初步治疗方案" not in obj:
            obj["初步治疗方案"] = obj.pop("治疗方案")
        return json.dumps(obj, ensure_ascii=False)
    return value


def _build_special_cases_sheet(doc_d1: pd.DataFrame, doc_d2: pd.DataFrame, special_cases: set[str]) -> pd.DataFrame:
    if (doc_d1 is None or doc_d1.empty) and (doc_d2 is None or doc_d2.empty):
        return pd.DataFrame()
    if not special_cases:
        return pd.DataFrame()
    # Template columns: prefer D2 decision structure if available.
    template_cols = list(doc_d2.columns) if doc_d2 is not None and not doc_d2.empty else list(doc_d1.columns)
    if not template_cols:
        return pd.DataFrame()

    rename_map = {"医生决策_原始JSON_治疗方案": "医生决策_原始JSON_初步治疗方案"}
    template_cols = [rename_map.get(str(c), c) for c in template_cols]

    rows: list[dict] = []
    for case_id in sorted(special_cases):
        row = None
        source = None
        if doc_d1 is not None and not doc_d1.empty:
            cid_col = doc_d1.columns[0]
            sel = doc_d1.loc[doc_d1[cid_col].astype(str) == str(case_id)]
            if not sel.empty:
                row = sel.iloc[0]
                source = "d1"
        if (row is None or not _row_has_output_for_stage("D1_Outpatient_Decision", doc_d1, row)) and doc_d2 is not None and not doc_d2.empty:
            cid_col = doc_d2.columns[0]
            sel = doc_d2.loc[doc_d2[cid_col].astype(str) == str(case_id)]
            if not sel.empty:
                row = sel.iloc[0]
                source = "d2"
        if row is None:
            continue
        out = {c: pd.NA for c in template_cols}
        out[template_cols[0]] = case_id
        if len(template_cols) > 1:
            out[template_cols[1]] = "终止于门诊决策（特殊情况）"

        # Copy same-name columns if present
        for c in template_cols[2:]:
            src_col = c
            if c == "医生决策_原始JSON_初步治疗方案":
                if source == "d1" and "医生决策_原始JSON_治疗方案" in doc_d1.columns:
                    src_col = "医生决策_原始JSON_治疗方案"
                elif source == "d2" and "医生决策_原始JSON_治疗方案" in doc_d2.columns:
                    src_col = "医生决策_原始JSON_治疗方案"
            if source == "d1" and doc_d1 is not None and src_col in doc_d1.columns and not is_empty(row.get(src_col)):
                out[c] = row.get(src_col)
            if source == "d2" and doc_d2 is not None and src_col in doc_d2.columns and not is_empty(row.get(src_col)):
                out[c] = row.get(src_col)

        # Map treatment plan into D2 key if needed
        if "医生决策_原始JSON_初步治疗方案" in template_cols:
            if is_empty(out.get("医生决策_原始JSON_初步治疗方案")):
                if source == "d1" and doc_d1 is not None:
                    if "医生决策_原始JSON_初步治疗方案" in doc_d1.columns and not is_empty(
                        row.get("医生决策_原始JSON_初步治疗方案")
                    ):
                        out["医生决策_原始JSON_初步治疗方案"] = row.get("医生决策_原始JSON_初步治疗方案")
                    elif "医生决策_原始JSON_治疗方案" in doc_d1.columns and not is_empty(row.get("医生决策_原始JSON_治疗方案")):
                        out["医生决策_原始JSON_初步治疗方案"] = row.get("医生决策_原始JSON_治疗方案")
                if source == "d2" and doc_d2 is not None:
                    if "医生决策_原始JSON_初步治疗方案" in doc_d2.columns and not is_empty(
                        row.get("医生决策_原始JSON_初步治疗方案")
                    ):
                        out["医生决策_原始JSON_初步治疗方案"] = row.get("医生决策_原始JSON_初步治疗方案")
                    elif "医生决策_原始JSON_治疗方案" in doc_d2.columns and not is_empty(row.get("医生决策_原始JSON_治疗方案")):
                        out["医生决策_原始JSON_初步治疗方案"] = row.get("医生决策_原始JSON_治疗方案")

        if "医生决策_原始JSON" in out and not is_empty(out.get("医生决策_原始JSON")):
            out["医生决策_原始JSON"] = _rename_treatment_key_in_json(out.get("医生决策_原始JSON"))

        rows.append(out)

    return pd.DataFrame(rows, columns=template_cols)


def run_status_audit(
    project_root: Path,
    data_root: Path,
    work_root: Path,
    tag: str,
    centers: list[str] | None,
    models: list[str] | None,
) -> None:
    data_root = (project_root / data_root).resolve()
    work_root = (project_root / work_root).resolve()

    run_id = _make_run_id(tag)
    work_dir = work_root / run_id
    audit_dir = work_dir / "status_audit"
    fixed_root = work_dir / "fixed_excels"
    if run_id == "latest":
        # Keep history: archive instead of deleting.
        def _unique_dest(base: Path) -> Path:
            if not base.exists():
                return base
            for i in range(1, 1000):
                cand = base.parent / f"{base.name}-{i}"
                if not cand.exists():
                    return cand
            raise RuntimeError(f"Cannot find unique archive dest for {base}")

        ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        archive_root = work_root / "archive"
        archive_root.mkdir(parents=True, exist_ok=True)
        archive_dir = _unique_dest(archive_root / f"{ts}_latest_status_audit")
        archive_dir.mkdir(parents=True, exist_ok=True)

        if audit_dir.exists():
            shutil.move(str(audit_dir), str(archive_dir / "status_audit"))
        if fixed_root.exists():
            shutil.move(str(fixed_root), str(archive_dir / "fixed_excels"))
    audit_dir.mkdir(parents=True, exist_ok=True)
    fixed_root.mkdir(parents=True, exist_ok=True)

    discovery = discover_datasets(data_root, centers=centers, models=models)
    (audit_dir / "discovery_warnings.json").write_text(json.dumps(discovery.warnings, ensure_ascii=False, indent=2), encoding="utf-8")

    diffs: list[StatusDiffRow] = []
    inferred_rows: list[pd.DataFrame] = []

    for ds in discovery.datasets:
        # Detect D1 decision special cases (based on content, not raw status)
        special_cases: set[str] = set()
        misplaced_cases: set[str] = set()
        doc_d1 = pd.DataFrame()
        doc_d2 = pd.DataFrame()
        try:
            doc_d2 = pd.read_excel(ds.doc_path, sheet_name="D2_Admission_Decision", engine="openpyxl")
            doc_d2, _ = canonicalize_sheet_columns(doc_d2)
        except Exception:
            doc_d2 = pd.DataFrame()

        # Detect misplaced D1 rows that were accidentally stored in D2 decision.
        try:
            doc_d1 = pd.read_excel(ds.doc_path, sheet_name="D1_Outpatient_Decision", engine="openpyxl")
            doc_d1, _ = canonicalize_sheet_columns(doc_d1)
            doc_fixed, _ = repair_doc_decision_nested_fields({"D1_Outpatient_Decision": doc_d1})
            doc_d1 = doc_fixed.get("D1_Outpatient_Decision", doc_d1)
            misplaced_cases = _detect_misplaced_d1_to_d2_cases(doc_d1, doc_d2)
        except Exception:
            misplaced_cases = set()

        if misplaced_cases:
            doc_d1, doc_d2 = _restore_misplaced_d1_from_d2_df(doc_d1, doc_d2, misplaced_cases)

        try:
            special_cases = detect_d1_decision_anomaly(doc_d1, doc_d2)
        except Exception:
            special_cases = set()

        # Prepare special-case export (D2 decision-style sheet)
        special_df = pd.DataFrame()
        case_overview_df = pd.DataFrame()
        if special_cases:
            special_df = _build_special_cases_sheet(doc_d1, doc_d2, special_cases)
            case_ids = sorted(special_cases, key=case_id_sort_key)
            case_overview_df = pd.DataFrame(
                [{"病例ID": cid, "模型名称": ds.model, "中心": ds.center} for cid in case_ids],
                columns=["病例ID", "模型名称", "中心"],
            )

        flow_df, inferred_status_map = _infer_status_maps_for_doc(ds.doc_path)
        if not flow_df.empty:
            flow_df.insert(0, "model", ds.model)
            flow_df.insert(0, "center", ds.center)
            inferred_rows.append(flow_df)

        for kind, src in [("doc", ds.doc_path), ("judge", ds.judge_path)]:
            try:
                rel = src.relative_to(data_root)
            except Exception:
                rel = Path("data") / ds.center / ("doc agent" if kind == "doc" else "judge agent") / src.name
            dst = fixed_root / rel
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dst)

            wb = load_workbook(dst)
            status_before = {sheet: _read_status_map(src, sheet) for sheet in SHEETS}
            changed = _apply_status_map_to_workbook(wb, inferred_status_map)
            if kind == "doc" and misplaced_cases:
                _restore_misplaced_d1_from_d2(wb, misplaced_cases)
            if kind == "doc" and special_cases and "D1_Outpatient_Decision" in wb.sheetnames:
                _apply_special_cases_to_doc_decision(wb["D1_Outpatient_Decision"], special_cases)
            if kind == "doc" and special_cases and "D2_Admission_Decision" in wb.sheetnames:
                _apply_special_cases_to_d2_decision(wb["D2_Admission_Decision"], special_cases)
            for sheet, case_id, old_status in changed:
                diffs.append(
                    StatusDiffRow(
                        center=ds.center,
                        model=ds.model,
                        kind=kind,
                        sheet=sheet,
                        case_id=case_id,
                        status_before=status_before.get(sheet, {}).get(case_id, old_status),
                        status_after=inferred_status_map.get(sheet, {}).get(case_id, ""),
                    )
                )
            wb.save(dst)

        # Write special-cases workbook (doc only; no judge sheet for D1 decision)
        if special_cases and not special_df.empty:
            special_root = work_dir / "special_cases" / ds.center / ds.model
            special_root.mkdir(parents=True, exist_ok=True)
            special_path = special_root / "D1决策_特殊情况转D2.xlsx"
            with pd.ExcelWriter(special_path, engine="openpyxl") as writer:
                if not case_overview_df.empty:
                    case_overview_df.to_excel(writer, sheet_name="Case_Overview", index=False)
                special_df.to_excel(writer, sheet_name="D2_Admission_Decision", index=False)

    # Write audit CSV
    csv_path = audit_dir / "status_diff.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["中心", "模型", "来源", "工作表", "阶段", "病例ID", "修改前状态", "修改后状态"])
        for d in diffs:
            w.writerow(
                [
                    d.center,
                    d.model,
                    _KIND_CN.get(d.kind, d.kind),
                    d.sheet,
                    _SHEET_CN.get(d.sheet, d.sheet),
                    d.case_id,
                    d.status_before,
                    d.status_after,
                ]
            )

    if inferred_rows:
        inferred = pd.concat(inferred_rows, ignore_index=True)
        inferred.to_csv(audit_dir / "flow_status_inferred.csv", index=False, encoding="utf-8-sig")
        issues = inferred.loc[inferred["issues"].astype(str).str.strip().ne(""), ["center", "model", "case_id", "flow_end_stage", "flow_end_detail", "issues"]].copy()
        if not issues.empty:
            issues.to_csv(audit_dir / "flow_status_issues.csv", index=False, encoding="utf-8-sig")

    # Rules.md
    (audit_dir / "rules.md").write_text(
        "\n".join(
            [
                "# Status Audit Rules",
                "",
                "- Source of truth: **doc sheet content** (non-empty extracted fields), not only column-B `状态`.",
                "- We infer `flow_end_stage/flow_end_detail` per case and write `flow_status_inferred.csv`.",
                "- Both doc/judge copies get updated column-B `状态` to inferred status (cross-sheet consistent).",
                "- Action: write **copies only** under `work/<run_id>/fixed_excels/` (never modify `data/`).",
                "- Also apply status conditional formatting (green=顺利通过, orange=终止, grey=未经过).",
                "",
            ]
        ),
        encoding="utf-8",
    )

    print(str(work_dir))
