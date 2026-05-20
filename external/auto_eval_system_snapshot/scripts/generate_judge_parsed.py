import pandas as pd
import json
import os
import logging
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

BASE_DIR = r"D:\研究生\项目\课题7-临床评测\自动测评系统\测评结果"
RED_FILL = PatternFill(start_color="FFB3B3", end_color="FFB3B3", fill_type="solid")
GRAY_FILL = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
GREEN_FILL = PatternFill(start_color="B3FFB3", end_color="B3FFB3", fill_type="solid")

SHEET_ORDER = ["D1_Outpatient_Loop", "D1_Outpatient_Decision", "D2_Admission_Loop", "D2_Admission_Decision", "D3_Surgery_Decision", "D4_Rehab_Plan"]
SHEET_NAMES_CN = {"D1_Outpatient_Loop": "门诊循环", "D1_Outpatient_Decision": "门诊决策", "D2_Admission_Loop": "入院循环", "D2_Admission_Decision": "入院决策", "D3_Surgery_Decision": "手术决策", "D4_Rehab_Plan": "康复计划"}

def safe_json_loads(val):
    if pd.isna(val) or str(val).strip() == "":
        return None, True
    try:
        parsed = json.loads(str(val))
        if isinstance(parsed, dict) and "Error" in parsed and len(parsed) == 1:
             return parsed, True
        return parsed, False
    except (json.JSONDecodeError, TypeError):
        return None, True

def has_data_in_any_future_sheet(case_id, current_sheet_idx, all_original_dfs):
    """检查该病例在所有后续sheet中是否有任何判官JSON数据"""
    for future_idx in range(current_sheet_idx + 1, len(SHEET_ORDER)):
        future_sheet = SHEET_ORDER[future_idx]
        if future_sheet not in all_original_dfs:
            continue
        future_df = all_original_dfs[future_sheet]
        id_col = "病例ID" if "病例ID" in future_df.columns else "CaseID"
        if id_col not in future_df.columns:
            continue
        case_rows = future_df[future_df[id_col] == case_id]
        if case_rows.empty:
            continue
        for col in case_rows.columns:
            if any(kw in col for kw in ['Judge', '判官', 'Gate']) and 'JSON' in col:
                val = case_rows[col].iloc[0]
                if pd.notna(val) and str(val).strip() != "" and str(val).strip().lower() != "nan":
                    return True
    return False

def determine_status_judge(row, sheet_name, sheet_idx, all_original_dfs):
    case_id = row.get('病例ID', None)
    judge_json_cols = [c for c in row.index if any(kw in c for kw in ['Judge', '判官', 'Gate']) and 'JSON' in c and not c.endswith('_JSON_')]
    
    current_has_data = False
    for col in judge_json_cols:
        val = row[col]
        if pd.notna(val) and str(val).strip() != "" and str(val).strip().lower() != "nan":
            current_has_data = True
            break
    
    if not current_has_data:
        return "未经过", GRAY_FILL
    
    has_future_data = has_data_in_any_future_sheet(case_id, sheet_idx, all_original_dfs)
    if has_future_data:
        return "顺利通过", GREEN_FILL
    else:
        if sheet_idx == len(SHEET_ORDER) - 1:
            return "顺利通过（流程完成）", GREEN_FILL
        return f"终止于{SHEET_NAMES_CN.get(sheet_name, sheet_name)}", RED_FILL

def process_excel(file_path):
    try:
        logger.info(f"处理: {file_path}")
        xls = pd.ExcelFile(file_path)
        all_original_dfs = {sheet: pd.read_excel(file_path, sheet_name=sheet) for sheet in xls.sheet_names}
        output_dfs = {}
        sheet_status_data = {}
        
        for sheet_idx, sheet_name in enumerate(SHEET_ORDER):
            if sheet_name not in all_original_dfs:
                continue
            df = all_original_dfs[sheet_name]
            id_col = "病例ID" if "病例ID" in df.columns else ("CaseID" if "CaseID" in df.columns else None)
            all_json_cols = [c for c in df.columns if "原始JSON" in c or "_原始JSON" in c]
            judge_json_cols = [c for c in all_json_cols if any(kw in c for kw in ["判官", "Judge", "Gate"])]
            
            if not judge_json_cols:
                continue
            logger.info(f"  Sheet: {sheet_name}")
            cols_to_keep = []
            if id_col: cols_to_keep.append(id_col)
            cols_to_keep.extend(judge_json_cols)
            subset_df = df[cols_to_keep].copy()
            final_concat_list = [subset_df]
            
            for j_col in judge_json_cols:
                parsed_results = subset_df[j_col].apply(safe_json_loads)
                data_list = [x[0] for x in parsed_results]
                valid_data = [d if d is not None else {} for d in data_list]
                flattened = pd.json_normalize(valid_data, sep='_')
                flattened.columns = [f"{j_col}_{c}" for c in flattened.columns]
                flattened.index = subset_df.index
                final_concat_list.append(flattened)
            
            final_df = pd.concat(final_concat_list, axis=1)
            status_col_data = []
            status_info = []
            for idx, row in final_df.iterrows():
                status_text, fill_color = determine_status_judge(row, sheet_name, sheet_idx, all_original_dfs)
                status_col_data.append(status_text)
                status_info.append((idx, status_text, fill_color))
            final_df.insert(1, '状态', status_col_data)
            output_dfs[sheet_name] = final_df
            sheet_status_data[sheet_name] = status_info
        
        output_path = file_path.replace("_CN.xlsx", "_CN_Judge_Parsed.xlsx")
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            for s_name, d_frame in output_dfs.items():
                d_frame.to_excel(writer, sheet_name=s_name, index=False)
        wb = load_workbook(output_path)
        for s_name, status_list in sheet_status_data.items():
            if s_name in wb.sheetnames:
                ws = wb[s_name]
                for idx, status_text, fill_color in status_list:
                    row_num = idx + 2
                    for cell in ws[row_num]:
                        cell.fill = fill_color
        wb.save(output_path)
        logger.info(f"  生成: {output_path}")
        return True
    except Exception as e:
        logger.error(f"失败 {file_path}: {e}")
        import traceback
        traceback.print_exc()
        return False

def main():
    logger.info("开始生成Judge Parsed...")
    count = 0
    for root, dirs, files in os.walk(BASE_DIR):
        for file in files:
            if file.endswith("_CN.xlsx") and "Evaluation_Summary" in file and "_Parsed" not in file:
                full_path = os.path.join(root, file)
                if process_excel(full_path):
                    count += 1
    logger.info(f"✅ 完成! {count}个文件")

if __name__ == "__main__":
    main()
