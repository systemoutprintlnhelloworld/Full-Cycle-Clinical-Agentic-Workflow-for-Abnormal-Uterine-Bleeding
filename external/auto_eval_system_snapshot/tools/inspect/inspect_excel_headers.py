from openpyxl import load_workbook
import os

excel_path = "output/evaluation_gemini-2.5-pro.xlsx"
if not os.path.exists(excel_path):
    print(f"File not found: {excel_path}")
    exit(1)

wb = load_workbook(excel_path)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    headers = [cell.value for cell in ws[1]]
    print(f"Sheet: {sheet}")
    # Check for "原始请求JSON" or "AI_Raw_JSON"
    found_ai_cols = [h for h in headers if "原始" in str(h) or "Raw_JSON" in str(h)]
    found_gate_cols = [h for h in headers if "Gate原始" in str(h) or "Gate_Raw_JSON" in str(h)]
    
    print(f"  Headers ({len(headers)}): {headers}")
    print(f"  AI Raw Cols Found: {found_ai_cols}")
    print(f"  Gate Raw Cols Found: {found_gate_cols}")
    print(f"  Has AI Raw JSON: {bool(found_ai_cols)}")
    print(f"  Has Gate Raw JSON: {bool(found_gate_cols)}")
