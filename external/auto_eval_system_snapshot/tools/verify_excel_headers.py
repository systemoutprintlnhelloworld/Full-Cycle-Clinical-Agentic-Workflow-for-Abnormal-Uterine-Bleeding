import sys
import os
import shutil
from openpyxl import load_workbook

# Add project root to path
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from auto_eval_system.modules.logger import EvaluationLogger

def verify_headers():
    test_output_dir = "output_verification_test"
    if os.path.exists(test_output_dir):
        shutil.rmtree(test_output_dir)
    
    print(f"Initializing Logger in {test_output_dir}...")
    logger_mod = EvaluationLogger(output_dir=test_output_dir, model_name="verification_test")
    
    excel_path = os.path.join(test_output_dir, "evaluation_verification_test.xlsx")
    if not os.path.exists(excel_path):
        print("❌ Excel file not created.")
        return False
        
    wb = load_workbook(excel_path)
    
    # Check List
    checks = {
        "门诊决策": "Gate原始JSON",
        "入院决策": "Gate原始JSON",
        "手术决策": "Gate原始JSON",
        "出院康复": "Gate原始JSON"
    }
    
    all_pass = True
    
    for sheet_name, target_col in checks.items():
        if sheet_name not in wb.sheetnames:
            print(f"❌ Sheet '{sheet_name}' missing.")
            all_pass = False
            continue
            
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        
        if target_col in headers:
            print(f"✅ Sheet '{sheet_name}': Found '{target_col}'")
        else:
            print(f"❌ Sheet '{sheet_name}': Missing '{target_col}'")
            print(f"   Existing headers: {headers}")
            all_pass = False
            
    # Cleanup
    # shutil.rmtree(test_output_dir)
    return all_pass

if __name__ == "__main__":
    if verify_headers():
        print("\nAll header checks passed!")
    else:
        print("\nHeader checks FAILED.")
        sys.exit(1)
